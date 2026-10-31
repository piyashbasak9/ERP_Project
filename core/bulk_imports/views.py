from io import BytesIO

from django.core.exceptions import ValidationError
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook

from core.branch.models import Branch

from .models import ImportBatch
from .services import (
    HEADERS,
    approve_batch,
    confirm_batch,
    error_report_rows,
    upload_batch,
    validate_batch,
)


def _accessible_branches(request):
    return request.user.accessible_branch_ids(
        getattr(request, "active_branch_id", None)
    )


def _batch_for_user(request, pk):
    return get_object_or_404(
        ImportBatch.objects.filter(branch_id__in=_accessible_branches(request)),
        pk=pk,
    )


@require_GET
def import_home(request):
    branches = Branch.unscoped.filter(pk__in=_accessible_branches(request)).order_by(
        "name"
    )
    batches = ImportBatch.objects.filter(
        branch_id__in=_accessible_branches(request)
    ).select_related("branch", "created_by", "approved_by")
    return render(
        request,
        "bulk_imports/home.html",
        {"branches": branches, "batches": batches},
    )


@require_POST
def import_upload(request):
    branch_id = request.POST.get("branch")
    branch = Branch.unscoped.filter(
        pk=branch_id,
        pk__in=_accessible_branches(request),
        is_active=True,
    ).first()
    if branch is None:
        return JsonResponse(
            {"success": False, "message": "Select an accessible active branch."},
            status=400,
        )
    uploaded_file = request.FILES.get("file")
    if uploaded_file is None:
        return JsonResponse(
            {"success": False, "message": "Choose an .xlsx workbook."},
            status=400,
        )
    try:
        batch, created = upload_batch(
            uploaded_file=uploaded_file,
            branch=branch,
            import_type=request.POST.get("import_type", ""),
            user=request.user,
        )
    except ValidationError as exc:
        return JsonResponse(
            {"success": False, "errors": exc.messages},
            status=400,
        )
    return JsonResponse(
        {
            "success": True,
            "batch_id": batch.pk,
            "status": batch.status,
            "created": created,
        },
        status=201 if created else 200,
    )


@require_POST
def import_validate(request, pk):
    batch = _batch_for_user(request, pk)
    if batch.created_by_id != request.user.pk and not request.user.is_superuser:
        raise Http404
    try:
        errors = validate_batch(batch)
    except ValidationError as exc:
        return JsonResponse(
            {"success": False, "errors": exc.messages},
            status=400,
        )
    return JsonResponse(
        {
            "success": not errors,
            "status": batch.status,
            "row_count": batch.row_count,
            "error_count": len(errors),
            "error_report": f"/imports/errors/{batch.pk}/",
        }
    )


@require_GET
def import_errors(request, pk):
    batch = _batch_for_user(request, pk)
    if not batch.validation_errors:
        raise Http404
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Validation errors")
    sheet.append(("Row", "Field", "Error"))
    for row in error_report_rows(batch):
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="import-errors-{batch.pk}.xlsx"'
    )
    return response


@require_POST
def import_confirm(request, pk):
    batch = _batch_for_user(request, pk)
    if batch.created_by_id != request.user.pk and not request.user.is_superuser:
        raise Http404
    try:
        confirm_batch(batch=batch, user=request.user, request=request)
    except ValidationError as exc:
        return JsonResponse(
            {"success": False, "errors": exc.messages},
            status=400,
        )
    return JsonResponse({"success": True, "status": batch.status})


@require_POST
def import_approve(request, pk):
    batch = _batch_for_user(request, pk)
    try:
        approve_batch(batch=batch, user=request.user, request=request)
    except ValidationError as exc:
        return JsonResponse(
            {"success": False, "errors": exc.messages},
            status=400,
        )
    return JsonResponse({"success": True, "status": batch.status})


@require_GET
def import_template(request, import_type):
    headers = HEADERS.get(import_type)
    if headers is None:
        raise Http404
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Import")
    sheet.append(headers)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{import_type}-template.xlsx"'
    )
    return response
