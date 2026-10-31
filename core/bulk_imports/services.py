import hashlib
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.accounting.models import Account
from core.accounting.services import JournalService
from core.audit.services import log_action
from core.branch.models import Member
from core.savings.models import SavingsAccount, SavingsTransaction
from core.savings.services import SavingsService

from .models import ImportBatch
from .security import MAX_FILE_BYTES, safe_report_cell, validate_xlsx_archive

ROW_LIMIT = 5000
HEADERS = {
    ImportBatch.TYPE_MEMBERS: (
        "member_code",
        "first_name",
        "last_name",
        "phone",
        "email",
        "date_of_birth",
        "admission_date",
    ),
    ImportBatch.TYPE_SAVINGS_OPENING: (
        "account_number",
        "amount",
        "transaction_date",
    ),
    ImportBatch.TYPE_GL_OPENING: (
        "account_code",
        "debit",
        "credit",
        "journal_date",
        "narration",
    ),
}

def upload_batch(*, uploaded_file, branch, import_type, user):
    if import_type not in HEADERS:
        raise ValidationError("Select a supported import type.")
    content = uploaded_file.read(MAX_FILE_BYTES + 1)
    validate_xlsx_archive(content, uploaded_file.name)
    digest = hashlib.sha256(content).hexdigest()
    existing = ImportBatch.objects.filter(file_hash=digest).first()
    if existing:
        if existing.branch_id != branch.pk or existing.import_type != import_type:
            raise ValidationError("This file was already submitted for a different import.")
        return existing, False
    try:
        with transaction.atomic():
            batch = ImportBatch.objects.create(
                branch=branch,
                created_by=user,
                import_type=import_type,
                file_hash=digest,
                source_name=uploaded_file.name[:255],
                source_file=ContentFile(content, name="upload.xlsx"),
            )
    except IntegrityError:
        existing = ImportBatch.objects.filter(file_hash=digest).first()
        if existing is None:
            raise
        if existing.branch_id != branch.pk or existing.import_type != import_type:
            raise ValidationError("This file was already submitted for a different import.")
        return existing, False
    log_action(
        "Create",
        "ImportBatch",
        batch.pk,
        new_values={
            "import_type": import_type,
            "file_hash": digest,
            "status": batch.status,
        },
        branch_id=branch.pk,
    )
    return batch, True


def _read_rows(batch):
    with batch.source_file.open("rb") as source:
        content = source.read()
    validate_xlsx_archive(content, batch.source_name)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, KeyError, ValueError) as exc:
        raise ValidationError("The uploaded workbook could not be read.") from exc
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    expected_headers = HEADERS[batch.import_type]
    normalized_headers = tuple(
        str(value).strip().lower() if value is not None else ""
        for value in (header_row or ())
    )
    if normalized_headers != expected_headers:
        workbook.close()
        raise ValidationError(
            "Workbook columns do not match the downloadable template."
        )
    result = []
    for row_number, values in enumerate(rows, start=2):
        if row_number > ROW_LIMIT + 1:
            workbook.close()
            raise ValidationError(f"Workbook exceeds the {ROW_LIMIT} row limit.")
        if not any(value is not None for value in values):
            continue
        result.append((row_number, dict(zip(expected_headers, values, strict=False))))
    workbook.close()
    return result


def _as_date(value, field):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError(f"{field} must be an ISO date.") from exc
    raise ValueError(f"{field} is required and must be a date.")


def _as_decimal(value, field):
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(f"{field} must be a decimal amount.") from exc
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
        raise ValueError(f"{field} must be a non-negative amount with at most 2 decimals.")
    return amount


def validate_batch(batch):
    if batch.status != ImportBatch.STATUS_UPLOADED:
        raise ValidationError("Only uploaded batches can be validated.")
    rows = _read_rows(batch)
    if not rows:
        errors = [{"row": 0, "field": "", "message": "Workbook contains no data rows."}]
        batch.validation_errors = errors
        batch.row_count = 0
        batch.status = ImportBatch.STATUS_INVALID
        batch.save(update_fields=("validation_errors", "row_count", "status"))
        return errors
    errors = []
    seen = set()
    debit_total = Decimal("0.00")
    credit_total = Decimal("0.00")
    journal_date = None
    for row_number, values in rows:
        try:
            if batch.import_type == ImportBatch.TYPE_MEMBERS:
                member_code = str(values["member_code"] or "").strip()
                if not member_code:
                    raise ValueError("member_code is required.")
                if member_code in seen or Member.unscoped.filter(
                    member_code=member_code
                ).exists():
                    raise ValueError("member_code already exists.")
                seen.add(member_code)
                if not str(values["first_name"] or "").strip():
                    raise ValueError("first_name is required.")
                if not str(values["last_name"] or "").strip():
                    raise ValueError("last_name is required.")
                if values["email"]:
                    from django.core.validators import validate_email

                    validate_email(str(values["email"]).strip())
                for field in ("date_of_birth", "admission_date"):
                    if values[field]:
                        _as_date(values[field], field)
                member = Member(
                    branch_id=batch.branch_id,
                    member_code=member_code,
                    first_name=str(values["first_name"]).strip(),
                    last_name=str(values["last_name"]).strip(),
                    phone=str(values["phone"] or "").strip(),
                    email=str(values["email"] or "").strip(),
                    date_of_birth=_as_date(values["date_of_birth"], "date_of_birth")
                    if values["date_of_birth"]
                    else None,
                    admission_date=_as_date(values["admission_date"], "admission_date")
                    if values["admission_date"]
                    else timezone.localdate(),
                )
                member.full_clean()
            elif batch.import_type == ImportBatch.TYPE_SAVINGS_OPENING:
                account_number = str(values["account_number"] or "").strip()
                if account_number in seen:
                    raise ValueError("account_number appears more than once.")
                seen.add(account_number)
                account = SavingsAccount.unscoped.filter(
                    account_number=account_number,
                    branch_id=batch.branch_id,
                ).first()
                if account is None:
                    raise ValueError("Savings account is not in the selected branch.")
                if _as_decimal(values["amount"], "amount") <= 0:
                    raise ValueError("amount must be greater than zero.")
                _as_date(values["transaction_date"], "transaction_date")
            elif batch.import_type == ImportBatch.TYPE_GL_OPENING:
                account_code = str(values["account_code"] or "").strip()
                account = Account.objects.filter(
                    code=account_code,
                    is_active=True,
                    allows_direct_posting=True,
                ).first()
                if account is None:
                    raise ValueError("account_code must identify an active postable account.")
                debit = _as_decimal(values["debit"], "debit")
                credit = _as_decimal(values["credit"], "credit")
                if (debit > 0) == (credit > 0):
                    raise ValueError("Enter a positive debit or credit, but not both.")
                debit_total += debit
                credit_total += credit
                row_journal_date = _as_date(values["journal_date"], "journal_date")
                if journal_date is None:
                    journal_date = row_journal_date
                elif row_journal_date != journal_date:
                    raise ValueError("All opening GL rows must use the same journal date.")
        except (ValueError, ValidationError) as exc:
            message = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
            errors.append(
                {"row": row_number, "field": "", "message": message[:240]}
            )
    if (
        batch.import_type == ImportBatch.TYPE_GL_OPENING
        and debit_total != credit_total
    ):
        errors.append(
            {
                "row": 0,
                "field": "debit/credit",
                "message": "Opening journal debits and credits must balance to zero.",
            }
        )
    batch.validation_errors = errors
    batch.row_count = len(rows)
    batch.status = (
        ImportBatch.STATUS_INVALID if errors else ImportBatch.STATUS_VALIDATED
    )
    batch.save(update_fields=("validation_errors", "row_count", "status"))
    log_action(
        "Validate",
        "ImportBatch",
        batch.pk,
        new_values={"status": batch.status, "row_count": batch.row_count},
        branch_id=batch.branch_id,
    )
    return errors


def confirm_batch(*, batch, user, request):
    if batch.status != ImportBatch.STATUS_VALIDATED:
        raise ValidationError("Only a successfully validated batch can be confirmed.")
    if batch.import_type in {
        ImportBatch.TYPE_SAVINGS_OPENING,
        ImportBatch.TYPE_GL_OPENING,
    }:
        batch.status = ImportBatch.STATUS_AWAITING_APPROVAL
        batch.save(update_fields=("status",))
        log_action(
            "SubmitForApproval",
            "ImportBatch",
            batch.pk,
            new_values={"status": batch.status},
            branch_id=batch.branch_id,
        )
        return batch
    _execute_import(batch=batch, user=user, request=request)
    return batch


@transaction.atomic
def approve_batch(*, batch, user, request):
    if batch.status != ImportBatch.STATUS_AWAITING_APPROVAL:
        raise ValidationError("This batch is not awaiting approval.")
    if batch.created_by_id == user.pk:
        raise ValidationError("The maker cannot approve their own opening-balance import.")
    if batch.import_type not in {
        ImportBatch.TYPE_SAVINGS_OPENING,
        ImportBatch.TYPE_GL_OPENING,
    }:
        raise ValidationError("Only opening-balance batches use maker-checker approval.")
    _execute_import(batch=batch, user=user, request=request)
    batch.approved_by = user
    batch.save(update_fields=("approved_by",))
    log_action(
        "Approve",
        "ImportBatch",
        batch.pk,
        new_values={"approved_by_id": user.pk, "status": batch.status},
        branch_id=batch.branch_id,
    )
    return batch


@transaction.atomic
def _execute_import(*, batch, user, request):
    try:
        return _execute_import_rows(batch=batch, user=user, request=request)
    except IntegrityError as exc:
        raise ValidationError(
            "The import conflicts with data that changed after validation."
        ) from exc


@transaction.atomic
def _execute_import_rows(*, batch, user, request):
    if batch.status not in {
        ImportBatch.STATUS_VALIDATED,
        ImportBatch.STATUS_AWAITING_APPROVAL,
    }:
        raise ValidationError("This batch is not ready to import.")
    rows = _read_rows(batch)
    if batch.import_type == ImportBatch.TYPE_MEMBERS:
        for _, values in rows:
            member = Member(
                branch_id=batch.branch_id,
                member_code=str(values["member_code"]).strip(),
                first_name=str(values["first_name"]).strip(),
                last_name=str(values["last_name"]).strip(),
                phone=str(values["phone"] or "").strip(),
                email=str(values["email"] or "").strip(),
                date_of_birth=_as_date(values["date_of_birth"], "date_of_birth")
                if values["date_of_birth"]
                else None,
                admission_date=_as_date(values["admission_date"], "admission_date")
                if values["admission_date"]
                else timezone.localdate(),
                created_by=user,
            )
            member.full_clean()
            member.save()
            log_action(
                "Create",
                "Member",
                member.pk,
                new_values={"member_code": member.member_code},
                branch_id=batch.branch_id,
            )
    elif batch.import_type == ImportBatch.TYPE_SAVINGS_OPENING:
        service = SavingsService(request)
        for row_number, values in rows:
            account = SavingsAccount.unscoped.get(
                account_number=str(values["account_number"]).strip(),
                branch_id=batch.branch_id,
            )
            key = uuid.uuid5(
                uuid.UUID(batch.file_hash[:32]),
                f"savings-opening-{row_number}",
            )
            service.deposit(
                account_id=account.pk,
                amount=_as_decimal(values["amount"], "amount"),
                transaction_date=_as_date(
                    values["transaction_date"], "transaction_date"
                ),
                payment_channel=SavingsTransaction.CHANNEL_CASH,
                idempotency_key=key,
                narration=f"Opening balance import {batch.file_hash[:12]}",
            )
    elif batch.import_type == ImportBatch.TYPE_GL_OPENING:
        first_date = _as_date(rows[0][1]["journal_date"], "journal_date")
        service = JournalService(request)
        entry = service.create_entry(
            branch_id=batch.branch_id,
            journal_date=first_date,
            narration=f"GL opening balances import {batch.file_hash[:12]}",
            source_type="gl_opening_import",
            source_id=str(batch.pk),
            idempotency_key=uuid.uuid5(
                uuid.UUID(batch.file_hash[:32]), "gl-opening-journal"
            ),
        )
        if not entry._idempotent_replay:
            for _, values in rows:
                account = Account.objects.get(
                    code=str(values["account_code"]).strip(),
                    is_active=True,
                    allows_direct_posting=True,
                )
                if _as_date(values["journal_date"], "journal_date") != first_date:
                    raise ValidationError(
                        "All opening GL rows must use the same journal date."
                    )
                service.add_line(
                    entry_id=entry.pk,
                    account_id=account.pk,
                    debit=_as_decimal(values["debit"], "debit"),
                    credit=_as_decimal(values["credit"], "credit"),
                    narration=str(values["narration"] or "")[:255],
                )
            service.post(entry_id=entry.pk)
    batch.status = ImportBatch.STATUS_IMPORTED
    batch.imported_at = timezone.now()
    batch.save(update_fields=("status", "imported_at"))
    log_action(
        "Import",
        "ImportBatch",
        batch.pk,
        new_values={"status": batch.status, "row_count": batch.row_count},
        branch_id=batch.branch_id,
    )
    file_name = batch.source_file.name
    transaction.on_commit(lambda: default_storage.delete(file_name))


def error_report_rows(batch):
    return [
        (
            error.get("row", ""),
            safe_report_cell(error.get("field", "")),
            safe_report_cell(error.get("message", "")),
        )
        for error in batch.validation_errors
    ]
