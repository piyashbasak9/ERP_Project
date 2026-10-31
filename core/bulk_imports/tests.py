import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.accounting.models import (
    Account,
    AccountingPeriod,
    FiscalYear,
    JournalEntry,
    JournalLine,
    SystemAccountMapping,
)
from core.accounts.models import Permission, Role, RolePermission, UserRole
from core.branch.models import Branch, Member

from .models import ImportBatch
from .security import (
    MAX_FILE_BYTES,
    safe_report_cell,
    validate_xlsx_archive,
)
from .services import (
    HEADERS,
    approve_batch,
    confirm_batch,
    upload_batch,
    validate_batch,
)


def workbook_bytes(import_type, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS[import_type])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def uploaded(content, name="import.xlsx"):
    return SimpleUploadedFile(
        name,
        content,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )


class BulkImportServiceTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        media_override = override_settings(MEDIA_ROOT=self.temp_media.name)
        media_override.enable()
        self.addCleanup(media_override.disable)
        self.addCleanup(self.temp_media.cleanup)
        self.branch = Branch.objects.create(name="Import Branch", code="IM1")
        self.other_branch = Branch.objects.create(name="Other Import Branch", code="IM2")
        self.user = get_user_model().objects.create_user(
            username="import-maker",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=self.branch,
        )

    def make_batch(self, import_type, rows, *, branch=None):
        return upload_batch(
            uploaded_file=uploaded(workbook_bytes(import_type, rows)),
            branch=branch or self.branch,
            import_type=import_type,
            user=self.user,
        )[0]

    def test_validation_reports_row_level_errors(self):
        batch = self.make_batch(
            ImportBatch.TYPE_MEMBERS,
            [
                ("IMPORT-1", "", "Member", "", "", None, None),
                ("IMPORT-2", "Valid", "Member", "", "bad-email", None, None),
            ],
        )

        errors = validate_batch(batch)

        self.assertEqual(batch.status, ImportBatch.STATUS_INVALID)
        self.assertEqual([error["row"] for error in errors], [2, 3])
        self.assertEqual(batch.row_count, 2)

    def test_validated_import_rolls_back_every_row_on_late_failure(self):
        batch = self.make_batch(
            ImportBatch.TYPE_MEMBERS,
            [
                ("ROLLBACK-1", "First", "Member", "", "", None, None),
                ("ROLLBACK-2", "Second", "Member", "", "", None, None),
            ],
        )
        self.assertEqual(validate_batch(batch), [])
        Member.unscoped.create(
            branch=self.branch,
            member_code="ROLLBACK-2",
            first_name="Concurrent",
            last_name="Insert",
        )

        with self.assertRaisesMessage(
            ValidationError, "conflicts with data that changed"
        ):
            confirm_batch(batch=batch, user=self.user, request=None)

        self.assertFalse(
            Member.unscoped.filter(member_code="ROLLBACK-1").exists()
        )
        self.assertEqual(batch.status, ImportBatch.STATUS_VALIDATED)

    def test_upload_is_idempotent_by_file_hash(self):
        content = workbook_bytes(
            ImportBatch.TYPE_MEMBERS,
            [("IDEMP-1", "Same", "File", "", "", None, None)],
        )
        first, created = upload_batch(
            uploaded_file=uploaded(content),
            branch=self.branch,
            import_type=ImportBatch.TYPE_MEMBERS,
            user=self.user,
        )
        second, created_again = upload_batch(
            uploaded_file=uploaded(content),
            branch=self.branch,
            import_type=ImportBatch.TYPE_MEMBERS,
            user=self.user,
        )

        self.assertTrue(created)
        self.assertFalse(created_again)
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(ImportBatch.objects.count(), 1)

    def test_branch_isolation_rejects_foreign_savings_accounts(self):
        from core.savings.models import SavingsProduct
        from core.savings.services import SavingsService

        member = Member.unscoped.create(
            branch=self.other_branch,
            member_code="OTHER-MEMBER",
            first_name="Other",
            last_name="Member",
            status=Member.STATUS_ACTIVE,
        )
        superuser = get_user_model().objects.create_superuser(
            username="import-system-user",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
        )
        product = SavingsProduct.objects.create(
            code="IMPORT-SAV",
            name="Import product",
            savings_type=SavingsProduct.TYPE_GENERAL,
        )
        request = type("Request", (), {"user": superuser, "active_branch_id": None})()
        account = SavingsService(request).open_account(
            branch_id=self.other_branch.pk,
            product_id=product.pk,
            opened_at=timezone.localdate(),
            member_id=member.pk,
        )
        batch = self.make_batch(
            ImportBatch.TYPE_SAVINGS_OPENING,
            [(account.account_number, "25.00", timezone.localdate().isoformat())],
        )

        errors = validate_batch(batch)

        self.assertEqual(batch.status, ImportBatch.STATUS_INVALID)
        self.assertIn("selected branch", errors[0]["message"])

    def test_savings_opening_import_is_maker_checker_and_posts_balanced_journal(self):
        from core.savings.models import SavingsProduct, SavingsTransaction
        from core.savings.services import SavingsService

        today = timezone.localdate()
        fiscal_year = FiscalYear.objects.create(
            name=f"Savings Import FY {today.year}",
            start_date=date(today.year, 1, 1),
            end_date=date(today.year, 12, 31),
        )
        AccountingPeriod.objects.create(
            fiscal_year=fiscal_year,
            number=1,
            name="Opening period",
            start_date=fiscal_year.start_date,
            end_date=fiscal_year.end_date,
        )
        cash = Account.objects.create(
            code="SAV-IMP-CASH",
            name="Import cash",
            account_type=Account.ASSET,
            normal_balance=Account.DEBIT,
        )
        savings_ledger = Account.objects.create(
            code="SAV-IMP-LIABILITY",
            name="Import member savings",
            account_type=Account.LIABILITY,
            normal_balance=Account.CREDIT,
        )
        SystemAccountMapping.objects.create(key="cash_in_hand", account=cash)
        SystemAccountMapping.objects.create(
            key="member_savings", account=savings_ledger
        )
        product = SavingsProduct.objects.create(
            code="SAV-IMP-PRODUCT",
            name="Import savings product",
            savings_type=SavingsProduct.TYPE_GENERAL,
        )
        member = Member.unscoped.create(
            branch=self.branch,
            member_code="SAV-IMP-MEMBER",
            first_name="Savings",
            last_name="Member",
            status=Member.STATUS_ACTIVE,
        )
        request = type(
            "Request", (), {"user": self.user, "active_branch_id": None}
        )()
        account = SavingsService(request).open_account(
            branch_id=self.branch.pk,
            product_id=product.pk,
            opened_at=today,
            member_id=member.pk,
        )
        batch = self.make_batch(
            ImportBatch.TYPE_SAVINGS_OPENING,
            [(account.account_number, "25.00", today.isoformat())],
        )
        self.assertEqual(validate_batch(batch), [])
        confirm_batch(batch=batch, user=self.user, request=request)
        self.assertEqual(batch.status, ImportBatch.STATUS_AWAITING_APPROVAL)
        with self.assertRaises(ValidationError):
            approve_batch(batch=batch, user=self.user, request=request)

        checker = get_user_model().objects.create_user(
            username="savings-import-checker",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=self.branch,
        )
        checker_request = type(
            "Request", (), {"user": checker, "active_branch_id": None}
        )()
        approve_batch(batch=batch, user=checker, request=checker_request)

        account.refresh_from_db()
        transaction_record = SavingsTransaction.unscoped.get(
            account=account,
            narration__startswith="Opening balance import"
        )
        journal = transaction_record.journal_entry
        lines = list(JournalLine.objects.filter(entry=journal))
        self.assertEqual(account.balance, Decimal("25.00"))
        self.assertEqual(batch.status, ImportBatch.STATUS_IMPORTED)
        self.assertEqual(journal.status, JournalEntry.STATUS_POSTED)
        self.assertEqual(
            sum((line.debit - line.credit for line in lines), Decimal("0.00")),
            Decimal("0.00"),
        )

    def test_unbalanced_gl_opening_is_rejected_before_approval(self):
        today = timezone.localdate()
        account = Account.objects.create(
            code="IMBALANCE",
            name="Imbalance test",
            account_type=Account.ASSET,
            normal_balance=Account.DEBIT,
        )
        batch = self.make_batch(
            ImportBatch.TYPE_GL_OPENING,
            [(account.code, "25.00", "0.00", today.isoformat(), "opening")],
        )

        errors = validate_batch(batch)

        self.assertEqual(batch.status, ImportBatch.STATUS_INVALID)
        self.assertTrue(
            any("must balance to zero" in error["message"] for error in errors)
        )

    def test_gl_opening_import_posts_one_balanced_journal_after_checker_approval(self):
        today = timezone.localdate()
        fiscal_year = FiscalYear.objects.create(
            name=f"Import FY {today.year}",
            start_date=date(today.year, 1, 1),
            end_date=date(today.year, 12, 31),
        )
        period = AccountingPeriod.objects.create(
            fiscal_year=fiscal_year,
            number=1,
            name="Opening period",
            start_date=fiscal_year.start_date,
            end_date=fiscal_year.end_date,
        )
        del period
        cash = Account.objects.create(
            code="IMP-CASH",
            name="Import cash",
            account_type=Account.ASSET,
            normal_balance=Account.DEBIT,
        )
        equity = Account.objects.create(
            code="IMP-EQUITY",
            name="Import opening equity",
            account_type=Account.EQUITY,
            normal_balance=Account.CREDIT,
        )
        content_rows = [
            (cash.code, "125.00", "0.00", today.isoformat(), "Opening cash"),
            (equity.code, "0.00", "125.00", today.isoformat(), "Opening equity"),
        ]
        batch = self.make_batch(ImportBatch.TYPE_GL_OPENING, content_rows)
        self.assertEqual(validate_batch(batch), [])

        maker_request = type(
            "Request", (), {"user": self.user, "active_branch_id": None}
        )()
        confirm_batch(batch=batch, user=self.user, request=maker_request)
        checker = get_user_model().objects.create_user(
            username="import-checker",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=self.branch,
        )
        checker_request = type(
            "Request", (), {"user": checker, "active_branch_id": None}
        )()
        approve_batch(batch=batch, user=checker, request=checker_request)

        entry = JournalEntry.unscoped.get(source_type="gl_opening_import")
        lines = list(JournalLine.objects.filter(entry=entry))
        self.assertEqual(batch.status, ImportBatch.STATUS_IMPORTED)
        self.assertEqual(entry.status, JournalEntry.STATUS_POSTED)
        self.assertEqual(
            sum((line.debit for line in lines), Decimal("0.00")),
            sum((line.credit for line in lines), Decimal("0.00")),
        )
        self.assertEqual(
            sum((line.debit - line.credit for line in lines), Decimal("0.00")),
            Decimal("0.00"),
        )
        self.assertEqual(batch.created_by_id, self.user.pk)
        self.assertEqual(batch.approved_by_id, checker.pk)


class BulkImportSecurityTests(TestCase):
    def test_malicious_uploads_are_rejected(self):
        with self.assertRaises(ValidationError):
            validate_xlsx_archive(b"not a workbook", "wrong.xlsx")

        with self.assertRaises(ValidationError):
            validate_xlsx_archive(b"PK\x03\x04" + b"x" * MAX_FILE_BYTES, "large.xlsx")

        bomb = BytesIO()
        with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", b"0" * (1024 * 1024))
        with self.assertRaises(ValidationError):
            validate_xlsx_archive(bomb.getvalue(), "bomb.xlsx")

        macro = BytesIO()
        with ZipFile(macro, "w", ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", b"office")
            archive.writestr("xl/vbaProject.bin", b"macro")
        with self.assertRaises(ValidationError):
            validate_xlsx_archive(macro.getvalue(), "macros.xlsx")

    def test_row_limit_is_enforced(self):
        temp_media = tempfile.TemporaryDirectory()
        media_override = override_settings(MEDIA_ROOT=temp_media.name)
        media_override.enable()
        self.addCleanup(media_override.disable)
        self.addCleanup(temp_media.cleanup)
        branch = Branch.objects.create(name="Row Limit Branch", code="IR3")
        user = get_user_model().objects.create_user(
            username="import-row-limit-user",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=branch,
        )
        content = workbook_bytes(
            ImportBatch.TYPE_MEMBERS,
            [("ROW-LIMIT", "Limit", "Test", "", "", None, None)],
        )
        batch, _ = upload_batch(
            uploaded_file=uploaded(content),
            branch=branch,
            import_type=ImportBatch.TYPE_MEMBERS,
            user=user,
        )

        with patch(
            "core.bulk_imports.services.ROW_LIMIT",
            0,
        ), self.assertRaises(ValidationError):
            validate_batch(batch)

    def test_error_report_cells_are_neutralized(self):
        for formula in ("=1+1", "+1+1", "-1+1", "@SUM(A1:A2)"):
            self.assertTrue(safe_report_cell(formula).startswith("'"))

    def test_validation_error_download_neutralizes_formula_cells(self):
        temp_media = tempfile.TemporaryDirectory()
        media_override = override_settings(MEDIA_ROOT=temp_media.name)
        media_override.enable()
        self.addCleanup(media_override.disable)
        self.addCleanup(temp_media.cleanup)
        branch = Branch.objects.create(name="Report Branch", code="IR1")
        user = get_user_model().objects.create_user(
            username="import-report-user",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=branch,
        )
        batch, _ = upload_batch(
            uploaded_file=uploaded(
                workbook_bytes(ImportBatch.TYPE_MEMBERS, [])
            ),
            branch=branch,
            import_type=ImportBatch.TYPE_MEMBERS,
            user=user,
        )
        batch.validation_errors = [
            {"row": 2, "field": "=FORMULA", "message": "@external"}
        ]
        batch.save(update_fields=("validation_errors",))
        permission = Permission.objects.create(
            route_name="Module_Import_errors",
            url_pattern="/import-test/errors/",
        )
        role = Role.objects.create(name="import-error-role")
        RolePermission.objects.create(role=role, permission=permission)
        UserRole.objects.create(user=user, role=role)
        client = Client()
        client.force_login(user)

        response = client.get(
            reverse("bulk_imports:Module_Import_errors", args=(batch.pk,))
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        values = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(values[1], (2, "'=FORMULA", "'@external"))


class BulkImportRouteTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(name="Import Route Branch", code="IR2")
        self.user = get_user_model().objects.create_user(
            username="import-route-user",
            password="A-strong-test-password-123",  # nosec B106 - isolated test account
            branch=self.branch,
        )
        self.client = Client()
        self.client.force_login(self.user)

    def grant_routes(self, *route_names):
        role = Role.objects.create(name=f"import-route-role-{Role.objects.count()}")
        for route_name in route_names:
            permission, _ = Permission.objects.get_or_create(
                route_name=route_name,
                defaults={
                    "url_pattern": f"/import-test/{route_name}/",
                    "module": "Imports",
                    "action": "test",
                },
            )
            RolePermission.objects.create(role=role, permission=permission)
        UserRole.objects.update_or_create(user=self.user, defaults={"role": role})

    def test_every_import_route_requires_its_permission(self):
        routes = (
            ("GET", "bulk_imports:Module_Import_list", ()),
            ("POST", "bulk_imports:Module_Import_upload", ()),
            ("POST", "bulk_imports:Module_Import_validate", (1,)),
            ("GET", "bulk_imports:Module_Import_errors", (1,)),
            ("POST", "bulk_imports:Module_Import_confirm", (1,)),
            ("POST", "bulk_imports:Module_Import_approve", (1,)),
            ("GET", "bulk_imports:Module_Import_template", ("members",)),
        )
        for method, route_name, args in routes:
            with self.subTest(route=route_name):
                response = getattr(self.client, method.lower())(
                    reverse(route_name, args=args)
                )
                self.assertEqual(response.status_code, 403)

    def test_every_import_mutation_route_rejects_get(self):
        route_specs = (
            ("Module_Import_upload", "bulk_imports:Module_Import_upload", ()),
            ("Module_Import_validate", "bulk_imports:Module_Import_validate", (1,)),
            ("Module_Import_confirm", "bulk_imports:Module_Import_confirm", (1,)),
            ("Module_Import_approve", "bulk_imports:Module_Import_approve", (1,)),
        )
        self.grant_routes(*(spec[0] for spec in route_specs))
        for _, route_name, args in route_specs:
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name, args=args))
                self.assertEqual(response.status_code, 405)
