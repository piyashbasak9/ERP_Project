from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from tempfile import SpooledTemporaryFile

from django.core.exceptions import FieldError, PermissionDenied
from django.db.models import QuerySet, Sum
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_GET

from core.audit.services import log_action
from core.branch.models import Branch, Center, Staff
from core.configuration.services import get_organization_configuration

from .registry import REPORTS, REPORTS_BY_SLUG

MAX_SYNCHRONOUS_EXPORT_ROWS = 5000
EXCEL_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@dataclass(frozen=True)
class ReportFilters:
    branch_ids: frozenset
    branch_id: int | None
    center_id: int | None
    officer_id: int | None
    start_date: date | None
    end_date: date
    status: str


def _integer_filter(request, key):
    value = request.GET.get(key, "").strip()
    if not value:
        return None
    try:
        return int(value)
    except ValueError as exc:
        raise ValueError(f"{key} must be an integer.") from exc


def _date_filter(request, key):
    value = request.GET.get(key, "").strip()
    if not value:
        return None
    parsed = parse_date(value)
    if parsed is None:
        raise ValueError(f"{key} must use YYYY-MM-DD format.")
    return parsed


def _filters_for(request):
    accessible = request.user.accessible_branch_ids(request.active_branch_id)
    branch_id = _integer_filter(request, "branch")
    if branch_id is None:
        branch_id = _integer_filter(request, "branch_id")
    if branch_id is not None and branch_id not in accessible:
        raise PermissionDenied("You cannot access the selected branch.")
    branch_ids = frozenset({branch_id} if branch_id is not None else accessible)

    center_id = _integer_filter(request, "center")
    if center_id is not None and not Center.unscoped.filter(
        pk=center_id, branch_id__in=branch_ids, is_active=True
    ).exists():
        raise PermissionDenied("You cannot access the selected center.")
    officer_id = _integer_filter(request, "field_officer")
    if officer_id is not None and not Staff.unscoped.filter(
        pk=officer_id, branch_id__in=branch_ids, is_active=True
    ).exists():
        raise PermissionDenied("You cannot access the selected field officer.")

    start_date = _date_filter(request, "start_date")
    end_date = _date_filter(request, "end_date") or timezone.localdate()
    if start_date and start_date > end_date:
        raise ValueError("The start date must be on or before the end date.")
    return ReportFilters(
        branch_ids=branch_ids,
        branch_id=branch_id,
        center_id=center_id,
        officer_id=officer_id,
        start_date=start_date,
        end_date=end_date,
        status=request.GET.get("status", "").strip(),
    )


def _report_for_slug(report_slug):
    try:
        return REPORTS_BY_SLUG[report_slug]
    except KeyError as exc:
        raise PermissionDenied("This report is not registered.") from exc


def _queryset(request, report, filters):
    rows = report.queryset(request, filters)
    if isinstance(rows, QuerySet):
        return rows
    return list(rows)


def _column_value(row, column):
    accessor = column.accessor or column.key
    if isinstance(row, dict):
        if accessor in row:
            return row[accessor]
        value = row
        for part in accessor.split("__"):
            value = value.get(part) if isinstance(value, dict) else None
        return value
    value = row
    for part in accessor.split("__"):
        value = getattr(value, part, None)
        if value is None:
            return None
    return value


def _json_value(value):
    if isinstance(value, (date,)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _materialize_rows(rows, report):
    return [
        {
            column.key: _column_value(row, column)
            for column in report.columns
        }
        for row in rows
    ]


def _totals(rows, report, values=None):
    if isinstance(rows, QuerySet):
        expressions = {}
        for key in report.total_keys:
            column = next(column for column in report.columns if column.key == key)
            field_name = column.accessor or column.key
            if "__" not in field_name:
                expressions[f"_report_total_{key}"] = Sum(field_name)
        try:
            aggregate = rows.aggregate(**expressions) if expressions else {}
        except FieldError:
            aggregate = None
        if aggregate is not None:
            totals = {
                key: aggregate.get(f"_report_total_{key}") or Decimal(0)
                for key in report.total_keys
                if f"_report_total_{key}" in expressions
            }
            if len(totals) == len(report.total_keys):
                return totals

    materialized = values if values is not None else _materialize_rows(rows, report)
    return {
        key: sum(
            (Decimal(str(row.get(key) or 0)) for row in materialized),
            Decimal(0),
        )
        for key in report.total_keys
    }


def _row_count(rows):
    return rows.count() if isinstance(rows, QuerySet) else len(rows)


def _sorted_rows(rows, report, sort_key, direction):
    known_columns = {column.key: column for column in report.columns}
    column = known_columns.get(sort_key)
    if column is None:
        return rows
    field_name = column.accessor or column.key
    if isinstance(rows, QuerySet) and all(
        character.isalnum() or character in "_-" for character in field_name
    ):
        prefix = "-" if direction == "desc" else ""
        try:
            return rows.order_by(f"{prefix}{field_name}")
        except FieldError:
            return rows
    if isinstance(rows, list):
        return sorted(
            rows,
            key=lambda row: _column_value(row, column) or "",
            reverse=direction == "desc",
        )
    return rows


def _page_values(request, report, rows):
    preference = getattr(request, "user_preference", None)
    default_page_size = preference.default_page_size if preference else 25
    page_size = min(
        max(_integer_filter(request, "rows") or default_page_size, 1), 100
    )
    page = max(_integer_filter(request, "page") or 1, 1)
    direction = request.GET.get("sord", "asc").lower()
    rows = _sorted_rows(
        rows,
        report,
        request.GET.get("sidx", ""),
        direction if direction in {"asc", "desc"} else "asc",
    )
    count = _row_count(rows)
    if isinstance(rows, QuerySet):
        page_rows = list(rows[(page - 1) * page_size : page * page_size])
    else:
        page_rows = rows[(page - 1) * page_size : page * page_size]
    page_count = (count + page_size - 1) // page_size
    values = _materialize_rows(page_rows, report)
    return page, page_count, count, values


def _organization_name():
    organization = get_organization_configuration()
    if organization.pk and organization.legal_name:
        return organization.legal_name
    head_office = Branch.unscoped.filter(
        branch_type=Branch.BRANCH_TYPE_HEAD_OFFICE
    ).order_by("pk").first()
    if head_office:
        return head_office.name
    root = Branch.unscoped.filter(parent__isnull=True).order_by("pk").first()
    return root.name if root else "Organization"


def _organization_logo(organization):
    logo = organization.logo
    if (
        not logo
        or not logo.name.lower().endswith((".jpg", ".jpeg", ".png"))
        or not logo.storage.exists(logo.name)
    ):
        return None
    with logo.open("rb") as logo_file:
        return logo_file.read()


def _filter_labels(filters):
    branches = Branch.unscoped.filter(pk__in=filters.branch_ids).order_by("name")
    branch_label = (
        Branch.unscoped.filter(pk=filters.branch_id).values_list("name", flat=True).first()
        if filters.branch_id
        else "All accessible branches"
    )
    values = [
        f"Branch: {branch_label}",
        f"Center: {Center.unscoped.filter(pk=filters.center_id).values_list('name', flat=True).first() or 'All'}",
        f"Field officer: {Staff.unscoped.filter(pk=filters.officer_id).values_list('full_name', flat=True).first() or 'All'}",
        f"Date range: {filters.start_date.isoformat() if filters.start_date else 'Beginning'} to {filters.end_date.isoformat()}",
        f"Status: {filters.status or 'All'}",
    ]
    return "; ".join(values), list(branches)


def _report_context(request, report):
    filters = _filters_for(request)
    rows = _queryset(request, report, filters)
    totals = _totals(rows, report)
    filter_text, branches = _filter_labels(filters)
    page, pages, count, page_rows = _page_values(request, report, rows)
    totals["record_count"] = count
    centers = Center.unscoped.filter(
        branch_id__in=filters.branch_ids, is_active=True
    ).order_by("name")
    officers = Staff.unscoped.filter(
        branch_id__in=filters.branch_ids, is_active=True
    ).order_by("full_name")
    return {
        "report": report,
        "filters": filters,
        "rows": page_rows,
        "totals": {key: _json_value(value) for key, value in totals.items()},
        "total_rows": [
            {
                "label": next(
                    column.label
                    for column in report.columns
                    if column.key == key
                ),
                "value": _json_value(totals[key]),
            }
            for key in report.total_keys
        ],
        "page": page,
        "pages": pages,
        "records": count,
        "branches": branches,
        "centers": centers,
        "officers": officers,
        "filter_text": filter_text,
        "organization_name": _organization_name(),
        "default_page_size": (
            request.user_preference.default_page_size
            if getattr(request, "user_preference", None)
            else 25
        ),
        "columns_json": [
            {"key": column.key, "label": column.label}
            for column in report.columns
        ],
        "view_url": reverse(f"reports:{report.route_name}"),
        "grid_url": reverse(f"reports:{report.grid_route_name}"),
        "excel_url": reverse(f"reports:{report.excel_route_name}"),
        "pdf_url": reverse(f"reports:{report.pdf_route_name}"),
    }


def _handle_report(request, report_slug, *, data=False):
    report = _report_for_slug(report_slug)
    context = _report_context(request, report)
    if data:
        return JsonResponse(
            {
                "page": context["page"],
                "total": context["pages"],
                "records": context["records"],
                "rows": [
                    {
                        key: _json_value(value)
                        for key, value in row.items()
                    }
                    for row in context["rows"]
                ],
                "userdata": context["totals"],
            }
        )
    return render(request, "reports/report.html", context)


@require_GET
def report_index(request):
    if request.user.is_superuser:
        available_routes = {report.route_name for report in REPORTS}
    else:
        try:
            role = request.user.user_role.role
        except AttributeError:
            available_routes = set()
        else:
            available_routes = set(
                role.permissions.filter(
                    is_allowed=True, is_active=True
                ).values_list("route_name", flat=True)
            )
    reports = [
        report for report in REPORTS if report.route_name in available_routes
    ]
    return render(
        request,
        "reports/index.html",
        {
            "reports": [
                {
                    "definition": report,
                    "url": reverse(f"reports:{report.route_name}"),
                }
                for report in reports
            ],
        },
    )


@require_GET
def report_view(request, report_slug):
    try:
        return _handle_report(request, report_slug)
    except ValueError as exc:
        return HttpResponse(str(exc), status=400)


@require_GET
def report_grid(request, report_slug):
    try:
        return _handle_report(request, report_slug, data=True)
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)


def _export_data(request, report_slug):
    report = _report_for_slug(report_slug)
    filters = _filters_for(request)
    rows = _queryset(request, report, filters)
    count = _row_count(rows)
    if count > MAX_SYNCHRONOUS_EXPORT_ROWS:
        return report, filters, rows, HttpResponse(
            f"This export contains {count:,} rows; the synchronous export limit is "
            f"{MAX_SYNCHRONOUS_EXPORT_ROWS:,}. Narrow the filters and try again.",
            status=413,
            content_type="text/plain; charset=utf-8",
        )
    matrix = _materialize_rows(rows, report)
    totals = _totals(rows, report, matrix)
    totals["record_count"] = count
    filter_text, _ = _filter_labels(filters)
    return report, filters, matrix, (totals, filter_text)


def _audit_export(
    request, report, filters, export_format, count, outcome="generated"
):
    log_action(
        "Export",
        "Report",
        report.slug,
        new_values={
            "format": export_format,
            "row_count": count,
            "outcome": outcome,
            "branch_id": filters.branch_id,
            "center_id": filters.center_id,
            "field_officer_id": filters.officer_id,
            "start_date": filters.start_date.isoformat() if filters.start_date else None,
            "end_date": filters.end_date.isoformat(),
            "status": filters.status,
        },
        branch_id=filters.branch_id,
    )


def _excel_file(workbook):
    spool = SpooledTemporaryFile(max_size=1_048_576, mode="w+b")  # noqa: SIM115 - FileResponse owns the stream.
    try:
        workbook.save(spool)
    except Exception:
        spool.close()
        raise
    spool.seek(0)
    return spool


@require_GET
def report_excel(request, report_slug):
    try:
        report, filters, rows, export_context = _export_data(request, report_slug)
    except ValueError as exc:
        return HttpResponse(str(exc), status=400)
    if isinstance(export_context, HttpResponse):
        _audit_export(
            request,
            report,
            filters,
            "xlsx",
            _row_count(rows),
            outcome="capped",
        )
        return export_context

    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.utils import get_column_letter

    totals, filter_text = export_context
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=report.title[:31])
    organization = get_organization_configuration()
    logo_data = _organization_logo(organization)
    if logo_data:
        logo = ExcelImage(BytesIO(logo_data))
        scale = min(120 / logo.width, 60 / logo.height, 1)
        logo.width *= scale
        logo.height *= scale
        sheet.add_image(logo, f"{get_column_letter(len(report.columns) + 2)}1")
    sheet.append([_organization_name()])
    sheet.append([report.title])
    sheet.append([filter_text])
    sheet.append([column.label for column in report.columns])
    for row in rows:
        sheet.append([_excel_value(row.get(column.key)) for column in report.columns])
    sheet.append([f"Total records: {len(rows)}"])
    if report.total_keys:
        sheet.append(
            [
                "Totals" if index == 0 else totals.get(column.key, "")
                for index, column in enumerate(report.columns)
            ]
        )
    excel_file = _excel_file(workbook)
    _audit_export(request, report, filters, "xlsx", len(rows))
    response = FileResponse(
        excel_file,
        content_type=EXCEL_CONTENT_TYPE,
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{report.slug}-{filters.end_date.isoformat()}.xlsx"'
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _excel_value(value):
    if isinstance(value, (date, Decimal, int, float, bool)) or value is None:
        return value
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(
        ("=", "+", "-", "@")
    ):
        return f"'{value}"
    return str(value)


@require_GET
def report_pdf(request, report_slug):
    try:
        report, filters, rows, export_context = _export_data(request, report_slug)
    except ValueError as exc:
        return HttpResponse(str(exc), status=400)
    if isinstance(export_context, HttpResponse):
        _audit_export(
            request,
            report,
            filters,
            "pdf",
            _row_count(rows),
            outcome="capped",
        )
        return export_context

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        Paragraph,
        SimpleDocTemplate,
        Table,
        TableStyle,
    )

    totals, filter_text = export_context
    from io import BytesIO

    buffer = BytesIO()
    story = []
    organization = get_organization_configuration()
    logo_data = _organization_logo(organization)
    if logo_data:
        logo = Image(
            BytesIO(logo_data),
            width=0.65 * inch,
            height=0.65 * inch,
            kind="proportional",
        )
        logo.hAlign = "LEFT"
        story.append(logo)
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.35 * inch,
        rightMargin=0.35 * inch,
    )
    styles = getSampleStyleSheet()
    data = [[Paragraph(_pdf_value(column.label), styles["BodyText"]) for column in report.columns]]
    for row in rows:
        data.append(
            [
                Paragraph(_pdf_value(row.get(column.key)), styles["BodyText"])
                for column in report.columns
            ]
        )
    if report.total_keys:
        data.append(
            [
                "Totals" if index == 0 else _pdf_value(totals.get(column.key, ""))
                for index, column in enumerate(report.columns)
            ]
        )
    data.append([f"Total records: {len(rows)}"])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183153")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.extend(
        [
            Paragraph(_pdf_value(_organization_name()), styles["Title"]),
            Paragraph(_pdf_value(report.title), styles["Heading2"]),
            Paragraph(_pdf_value(filter_text), styles["BodyText"]),
            table,
        ]
    )
    document.build(story)
    _audit_export(request, report, filters, "pdf", len(rows))
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="{report.slug}-{filters.end_date.isoformat()}.pdf"'
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _pdf_value(value):
    return str(value if value is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def report_route_patterns():
    from django.urls import path

    routes = []
    for report in REPORTS:
        routes.extend(
            [
                path(
                    f"{report.slug}/",
                    report_view,
                    {"report_slug": report.slug},
                    name=report.route_name,
                ),
                path(
                    f"{report.slug}/data/",
                    report_grid,
                    {"report_slug": report.slug},
                    name=report.grid_route_name,
                ),
                path(
                    f"{report.slug}/export/excel/",
                    report_excel,
                    {"report_slug": report.slug},
                    name=report.excel_route_name,
                ),
                path(
                    f"{report.slug}/export/pdf/",
                    report_pdf,
                    {"report_slug": report.slug},
                    name=report.pdf_route_name,
                ),
            ]
        )
    return routes
