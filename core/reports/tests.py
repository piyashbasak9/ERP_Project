import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import Sum
from django.test import RequestFactory
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image as PillowImage

from core.accounting.models import (
    Account,
    AccountingPeriod,
    FiscalYear,
    JournalEntry,
    JournalLine,
    SystemAccountMapping,
)
from core.accounting.services import JournalService
from core.accounts.models import Permission, Role, RolePermission, UserRole
from core.branch.models import Branch, Center, Member, Staff, UserBranchAccess
from core.collections.models import CollectionSheet, CollectionSheetItem
from core.configuration.models import OrganizationConfiguration
from core.loans.models import (
    Loan,
    LoanApplication,
    LoanInstallment,
    LoanPARBucket,
    LoanPayment,
    LoanProduct,
)
from core.savings.models import SavingsAccount, SavingsProduct

from .registry import REPORTS, REPORTS_BY_SLUG


def _route_names(report):
    return (
        report.route_name,
        report.grid_route_name,
        report.excel_route_name,
        report.pdf_route_name,
    )


def _create_user(branch):
    return get_user_model().objects.create_user(
        username=f"report-user-{uuid.uuid4().hex[:12]}",
        password="A-strong-test-password-2026",
        branch=branch,
    )


def _grant_routes(user, route_names):
    role = Role.objects.create(name=f"report-role-{uuid.uuid4().hex[:12]}")
    UserRole.objects.create(user=user, role=role)
    for route_name in route_names:
        permission, _ = Permission.objects.get_or_create(
            route_name=route_name,
            defaults={
                "url_pattern": f"/report-test/{route_name}/",
                "module": "Reports",
                "action": "view",
                "description": "Report test route",
            },
        )
        RolePermission.objects.create(role=role, permission=permission)
    return role


def _grant_report(user, report):
    return _grant_routes(user, _route_names(report))


@pytest.mark.django_db
@pytest.mark.parametrize("report", REPORTS, ids=lambda report: report.slug)
def test_every_registered_report_requires_its_own_route_permissions(client, report):
    branch = Branch.objects.create(name="Report permission branch", code="RPT-PERM")
    user = _create_user(branch)
    all_names = {
        route_name
        for definition in REPORTS
        for route_name in _route_names(definition)
    }
    _grant_routes(user, all_names - set(_route_names(report)))
    client.force_login(user)

    for route_name in _route_names(report):
        response = client.get(reverse(f"reports:{route_name}"))
        assert response.status_code == 403, route_name


@pytest.mark.django_db
@pytest.mark.parametrize("report", REPORTS, ids=lambda report: report.slug)
def test_every_registered_report_rejects_forged_branch_ids_on_page_grid_and_exports(
    client, report
):
    branch = Branch.objects.create(name="Report allowed branch", code="RPT-A")
    forbidden_branch = Branch.objects.create(
        name="Report forbidden branch", code="RPT-B"
    )
    user = _create_user(branch)
    _grant_report(user, report)
    client.force_login(user)

    page = client.get(
        reverse(f"reports:{report.route_name}"),
        {"branch": forbidden_branch.pk},
    )
    grid = client.get(
        reverse(f"reports:{report.grid_route_name}"),
        {"branch": forbidden_branch.pk},
    )
    excel = client.get(
        reverse(f"reports:{report.excel_route_name}"),
        {"branch": forbidden_branch.pk},
    )
    pdf = client.get(
        reverse(f"reports:{report.pdf_route_name}"),
        {"branch": forbidden_branch.pk},
    )

    assert page.status_code == 403
    assert grid.status_code == 403
    assert excel.status_code == 403
    assert pdf.status_code == 403


@pytest.mark.django_db
def test_user_access_review_is_branch_scoped_and_filters_user_status(client):
    branch = Branch.objects.create(name="Access review branch", code="ACCESS-A")
    other_branch = Branch.objects.create(
        name="Other access branch", code="ACCESS-B"
    )
    user = _create_user(branch)
    inactive_user = get_user_model().objects.create_user(
        username=f"inactive-access-{uuid.uuid4().hex[:8]}",
        password="A-strong-test-password-123",
        branch=branch,
        is_active=False,
    )
    outside_user = _create_user(other_branch)
    report = REPORTS_BY_SLUG["user-access-review"]
    _grant_report(user, report)
    outside_role = Role.objects.create(name=f"outside-role-{uuid.uuid4().hex[:8]}")
    UserRole.objects.create(user=outside_user, role=outside_role)
    inactive_role = Role.objects.create(
        name=f"inactive-role-{uuid.uuid4().hex[:8]}"
    )
    UserRole.objects.create(user=inactive_user, role=inactive_role)
    client.force_login(user)

    active_response = client.get(
        reverse(f"reports:{report.grid_route_name}"),
        {"status": "active"},
    )
    inactive_response = client.get(
        reverse(f"reports:{report.grid_route_name}"),
        {"status": "inactive"},
    )
    forged_branch_response = client.get(
        reverse(f"reports:{report.grid_route_name}"),
        {"branch": other_branch.pk},
    )

    assert active_response.status_code == 200
    active_rows = active_response.json()["rows"]
    assert all(row["branch"] == branch.name for row in active_rows)
    assert all(row["username"] != outside_user.username for row in active_rows)
    assert inactive_response.status_code == 200
    assert [row["username"] for row in inactive_response.json()["rows"]] == [
        inactive_user.username
    ]
    assert forged_branch_response.status_code == 403


@pytest.mark.django_db
def test_excel_and_pdf_exports_include_organization_logo(
    client, request, tmp_path, settings
):
    settings.MEDIA_ROOT = tmp_path
    cache.clear()
    request.addfinalizer(cache.clear)
    branch = Branch.objects.create(name="Branded report branch", code="RPT-BRAND")
    organization = OrganizationConfiguration.objects.create(
        legal_name="Branded Organization"
    )
    image = BytesIO()
    PillowImage.new("RGB", (20, 20), color="blue").save(image, format="PNG")
    organization.logo.save(
        "organization-logo.png",
        SimpleUploadedFile("organization-logo.png", image.getvalue(), "image/png"),
        save=True,
    )
    user = _create_user(branch)
    report = REPORTS_BY_SLUG["trial-balance"]
    _grant_report(user, report)
    client.force_login(user)

    excel = client.get(reverse(f"reports:{report.excel_route_name}"))
    workbook = load_workbook(BytesIO(b"".join(excel.streaming_content)))
    assert len(workbook.active._images) == 1
    workbook.close()

    pdf = client.get(reverse(f"reports:{report.pdf_route_name}"))
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")


@pytest.mark.django_db
@pytest.mark.parametrize("report", REPORTS, ids=lambda report: report.slug)
def test_every_registered_report_exports_audited_openpyxl_and_pdf_files(
    client, report
):
    from core.audit.models import AuditLog

    branch = Branch.objects.create(name="Report export branch", code="RPT-EXP")
    user = _create_user(branch)
    _grant_report(user, report)
    client.force_login(user)
    export_filters = {
        "branch": branch.pk,
        "start_date": date(2026, 1, 1).isoformat(),
        "end_date": timezone.localdate().isoformat(),
        "status": "posted",
    }

    excel_response = client.get(
        reverse(f"reports:{report.excel_route_name}"), export_filters
    )
    assert excel_response.status_code == 200
    excel_bytes = b"".join(excel_response.streaming_content)
    workbook = load_workbook(BytesIO(excel_bytes), read_only=True)
    exported_rows = list(workbook.active.iter_rows(values_only=True))
    assert exported_rows[0][0] == branch.name
    assert exported_rows[1][0] == report.title
    assert f"Branch: {branch.name}" in exported_rows[2][0]
    assert f"Date range: 2026-01-01 to {timezone.localdate().isoformat()}" in exported_rows[2][0]
    assert "Status: posted" in exported_rows[2][0]
    assert len(exported_rows) >= 4
    workbook.close()

    pdf_response = client.get(
        reverse(f"reports:{report.pdf_route_name}"), export_filters
    )
    assert pdf_response.status_code == 200
    assert pdf_response.content.startswith(b"%PDF-")
    assert pdf_response.content.rstrip().endswith(b"%%EOF")
    offset = int(pdf_response.content.rsplit(b"startxref\n", 1)[1].splitlines()[0])
    assert pdf_response.content[offset : offset + 4] == b"xref"

    audit_rows = AuditLog.objects.filter(
        entity_name="Report",
        entity_id=report.slug,
        action="Export",
    )
    assert audit_rows.count() == 2
    assert {
        row.new_values["format"] for row in audit_rows
    } == {"xlsx", "pdf"}


@pytest.mark.django_db
def test_report_index_lists_only_reports_the_role_can_open(client):
    branch = Branch.objects.create(name="Report index branch", code="RPT-IDX")
    user = _create_user(branch)
    selected, hidden = REPORTS[:2]
    _grant_routes(
        user,
        ("Module_Reports_index", selected.route_name),
    )
    client.force_login(user)

    response = client.get(reverse("reports:Module_Reports_index"))

    assert response.status_code == 200
    content = response.content.decode()
    assert selected.title in content
    assert hidden.title not in content


@pytest.mark.django_db
def test_large_synchronous_export_is_capped_with_audited_explanation(client):
    from core.audit.models import AuditLog

    branch = Branch.objects.create(name="Capped report branch", code="RPT-CAP")
    user = _create_user(branch)
    report = next(item for item in REPORTS if item.slug == "member-list")
    _grant_report(user, report)
    for index in range(2):
        Member.unscoped.create(
            branch=branch,
            member_code=f"RPT-CAP-{index}",
            first_name=f"Member{index}",
            last_name="Capped",
        )
    client.force_login(user)

    with patch("core.reports.views.MAX_SYNCHRONOUS_EXPORT_ROWS", 1):
        response = client.get(reverse(f"reports:{report.excel_route_name}"))

    assert response.status_code == 413
    assert b"Narrow the filters" in response.content
    audit = AuditLog.objects.get(entity_name="Report", entity_id=report.slug)
    assert audit.new_values["outcome"] == "capped"
    assert audit.new_values["row_count"] == 2


@pytest.mark.django_db
def test_member_report_filters_by_date_branch_and_status(client):
    branch_one = Branch.objects.create(name="Filter branch one", code="RPT-F1")
    branch_two = Branch.objects.create(name="Filter branch two", code="RPT-F2")
    user = _create_user(branch_one)
    UserBranchAccess.objects.create(user=user, branch=branch_two)
    report = next(item for item in REPORTS if item.slug == "member-list")
    _grant_report(user, report)
    today = timezone.localdate()
    Member.unscoped.create(
        branch=branch_one,
        member_code="RPT-FILTER-1",
        first_name="Active",
        last_name="One",
        admission_date=today,
        status=Member.STATUS_ACTIVE,
    )
    selected = Member.unscoped.create(
        branch=branch_two,
        member_code="RPT-FILTER-2",
        first_name="Inactive",
        last_name="Two",
        admission_date=today - timedelta(days=2),
        status=Member.STATUS_INACTIVE,
    )
    Member.unscoped.create(
        branch=branch_two,
        member_code="RPT-FILTER-3",
        first_name="Old",
        last_name="Member",
        admission_date=today - timedelta(days=90),
        status=Member.STATUS_ACTIVE,
    )
    client.force_login(user)

    response = client.get(
        reverse(f"reports:{report.grid_route_name}"),
        {
            "branch": branch_two.pk,
            "start_date": (today - timedelta(days=5)).isoformat(),
            "end_date": today.isoformat(),
            "status": Member.STATUS_INACTIVE,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["records"] == 1
    assert payload["rows"][0]["code"] == selected.member_code


@pytest.mark.django_db
def test_financial_reports_balance_and_match_posted_ledger(client):
    branch = Branch.objects.create(name="Financial report branch", code="RPT-FIN")
    user = _create_user(branch)
    report_slugs = {
        "general-ledger",
        "cash-book",
        "income-statement",
        "balance-sheet",
        "balance-summary",
        "receipts-payments",
    }
    reports = [report for report in REPORTS if report.slug in report_slugs]
    _grant_routes(
        user,
        tuple(
            route_name
            for report in reports
            for route_name in (report.route_name, report.grid_route_name)
        ),
    )
    today = timezone.localdate()
    fiscal_year = FiscalYear.objects.create(
        name="Financial report test year",
        start_date=date(today.year, 1, 1),
        end_date=date(today.year, 12, 31),
    )
    AccountingPeriod.objects.create(
        fiscal_year=fiscal_year,
        number=1,
        name="Financial report annual period",
        start_date=fiscal_year.start_date,
        end_date=fiscal_year.end_date,
    )
    cash = Account.objects.create(
        code="RPT-FIN-CASH",
        name="Report cash",
        account_type=Account.ASSET,
        normal_balance=Account.DEBIT,
    )
    savings_liability = Account.objects.create(
        code="RPT-FIN-SAV",
        name="Member savings",
        account_type=Account.LIABILITY,
        normal_balance=Account.CREDIT,
    )
    income = Account.objects.create(
        code="RPT-FIN-INC",
        name="Interest income",
        account_type=Account.INCOME,
        normal_balance=Account.CREDIT,
    )
    expense = Account.objects.create(
        code="RPT-FIN-EXP",
        name="Operating expense",
        account_type=Account.EXPENSE,
        normal_balance=Account.DEBIT,
    )
    SystemAccountMapping.objects.create(key="cash_in_hand", account=cash)
    request = RequestFactory().post("/reports/")
    request.user = user
    request.active_branch_id = None
    service = JournalService(request)

    def post_journal(narration, lines):
        entry = service.create_entry(
            branch_id=branch.pk,
            journal_date=today,
            narration=narration,
        )
        for account, debit, credit in lines:
            service.add_line(
                entry_id=entry.pk,
                account_id=account.pk,
                debit=Decimal(debit),
                credit=Decimal(credit),
            )
        service.post(entry_id=entry.pk)

    post_journal(
        "Member savings opening ledger",
        [(cash, "1250.00", "0.00"), (savings_liability, "0.00", "1250.00")],
    )
    post_journal(
        "Interest income",
        [(cash, "200.00", "0.00"), (income, "0.00", "200.00")],
    )
    post_journal(
        "Operating expense",
        [(expense, "50.00", "0.00"), (cash, "0.00", "50.00")],
    )
    member = Member.unscoped.create(
        branch=branch,
        member_code="RPT-FIN-MEMBER",
        first_name="Ledger",
        last_name="Member",
        status=Member.STATUS_ACTIVE,
    )
    product = SavingsProduct.objects.create(
        code="RPT-FIN-PRODUCT",
        name="Ledger savings product",
        savings_type=SavingsProduct.TYPE_GENERAL,
    )
    savings_account = SavingsAccount(
        branch=branch,
        product=product,
        member=member,
        account_number="RPT-FIN-SAV-0001",
        status=SavingsAccount.STATUS_ACTIVE,
        balance=Decimal("1250.00"),
        opened_at=today,
        created_by=user,
    )
    savings_account._savings_update_authorized = True
    savings_account.save()
    client.force_login(user)
    query = {
        "branch": branch.pk,
        "start_date": fiscal_year.start_date.isoformat(),
        "end_date": today.isoformat(),
    }

    responses = {
        report.slug: client.get(
            reverse(f"reports:{report.grid_route_name}"), query
        )
        for report in reports
    }
    payloads = {slug: response.json() for slug, response in responses.items()}
    assert all(response.status_code == 200 for response in responses.values())

    ledger = payloads["general-ledger"]
    assert ledger["userdata"]["debit"] == "1500.00"
    assert ledger["userdata"]["credit"] == "1500.00"
    cash_book = payloads["cash-book"]
    assert cash_book["userdata"]["receipt"] == "1450.00"
    assert cash_book["userdata"]["payment"] == "50.00"
    receipts_payments = payloads["receipts-payments"]
    assert receipts_payments["userdata"]["receipts"] == "1450.00"
    assert receipts_payments["userdata"]["payments"] == "50.00"
    income_statement = payloads["income-statement"]
    assert income_statement["userdata"]["amount"] == "150.00"
    balance_sheet_rows = payloads["balance-sheet"]["rows"]
    assets = sum(
        (Decimal(row["amount"]) for row in balance_sheet_rows if row["account_type"] == Account.ASSET),
        Decimal("0.00"),
    )
    liabilities = sum(
        (Decimal(row["amount"]) for row in balance_sheet_rows if row["account_type"] == Account.LIABILITY),
        Decimal("0.00"),
    )
    equity = sum(
        (Decimal(row["amount"]) for row in balance_sheet_rows if row["account_type"] == Account.EQUITY),
        Decimal("0.00"),
    )
    assert assets == Decimal("1400.00")
    assert liabilities + equity == assets

    summary = next(item for item in REPORTS if item.slug == "balance-summary")
    summary_payload = client.get(
        reverse(f"reports:{summary.grid_route_name}"), query
    ).json()
    ledger_balance = (
        JournalLine.objects.filter(
            account=savings_liability,
            entry__branch=branch,
            entry__status=JournalEntry.STATUS_POSTED,
        ).aggregate(
            debit=Sum("debit"),
            credit=Sum("credit"),
        )
    )
    assert summary_payload["userdata"]["balance"] == str(
        ledger_balance["credit"] - ledger_balance["debit"]
    )


@pytest.mark.django_db
def test_loan_reports_match_disbursement_schedule_overdue_and_writeoff_data(client):
    branch = Branch.objects.create(name="Loan report branch", code="RPT-LOAN")
    user = _create_user(branch)
    report_slugs = {
        "disbursement-register",
        "outstanding-loans",
        "overdue-loans",
        "aging",
        "write-offs",
        "repayment-schedule",
        "portfolio-at-risk",
        "loan-statement",
        "classification",
    }
    reports = [report for report in REPORTS if report.slug in report_slugs]
    _grant_routes(
        user,
        tuple(
            route_name
            for report in reports
            for route_name in (report.route_name, report.grid_route_name)
        ),
    )
    today = timezone.localdate()
    member = Member.unscoped.create(
        branch=branch,
        member_code="RPT-LOAN-MEMBER",
        first_name="Loan",
        last_name="Member",
        status=Member.STATUS_ACTIVE,
    )
    product = LoanProduct.objects.create(
        code="RPT-LOAN-PRODUCT",
        name="Report loan product",
        interest_method=LoanProduct.METHOD_FLAT,
        annual_interest_rate=0,
        min_amount=1,
        max_amount=10000,
        min_term=1,
        max_term=12,
        installment_frequency=LoanProduct.FREQUENCY_MONTHLY,
    )
    fiscal_year = FiscalYear.objects.create(
        name="Loan report fiscal year",
        start_date=date(today.year, 1, 1),
        end_date=date(today.year, 12, 31),
    )
    period = AccountingPeriod.objects.create(
        fiscal_year=fiscal_year,
        number=1,
        name="Loan report period",
        start_date=fiscal_year.start_date,
        end_date=fiscal_year.end_date,
    )
    journal = JournalEntry.unscoped.create(
        branch=branch,
        fiscal_year=fiscal_year,
        period=period,
        voucher_type="journal",
        voucher_number="RPT-LOAN-JOURNAL",
        journal_date=today,
        narration="Loan report fixture",
        created_by=user,
    )
    application = LoanApplication.unscoped.create(
        branch=branch,
        member=member,
        product=product,
        requested_amount=Decimal("1000.00"),
        approved_amount=Decimal("1000.00"),
        term=1,
        purpose="Report fixture",
        status=LoanApplication.STATUS_DISBURSED,
        created_by=user,
    )
    loan = Loan.unscoped.create(
        branch=branch,
        application=application,
        member=member,
        product=product,
        principal=Decimal("1000.00"),
        outstanding_principal=Decimal("700.00"),
        status=Loan.STATUS_ACTIVE,
        disbursed_at=today,
        disbursed_by=user,
        journal_entry=journal,
        idempotency_key=uuid.uuid4(),
    )
    LoanInstallment.objects.create(
        loan=loan,
        schedule_version=loan.schedule_version,
        number=1,
        due_date=today - timedelta(days=10),
        principal_due=Decimal("100.00"),
        interest_due=Decimal("10.00"),
        fee_due=Decimal("0.00"),
    )
    LoanPayment.unscoped.create(
        branch=branch,
        loan=loan,
        amount=Decimal("10.00"),
        principal_paid=Decimal("10.00"),
        payment_date=today,
        idempotency_key=uuid.uuid4(),
        journal_entry=journal,
        created_by=user,
    )
    writeoff_application = LoanApplication.unscoped.create(
        branch=branch,
        member=member,
        product=product,
        requested_amount=Decimal("300.00"),
        approved_amount=Decimal("300.00"),
        term=1,
        purpose="Write-off fixture",
        status=LoanApplication.STATUS_DISBURSED,
        created_by=user,
    )
    Loan.unscoped.create(
        branch=branch,
        application=writeoff_application,
        member=member,
        product=product,
        principal=Decimal("300.00"),
        outstanding_principal=Decimal("0.00"),
        written_off_amount=Decimal("300.00"),
        recovered_amount=Decimal("50.00"),
        status=Loan.STATUS_WRITTEN_OFF,
        disbursed_at=today,
        disbursed_by=user,
        journal_entry=journal,
        idempotency_key=uuid.uuid4(),
        write_off_reason="Unrecoverable account",
        closed_at=today,
    )
    LoanPARBucket.objects.create(
        name="1-30 days",
        minimum_days=1,
        maximum_days=30,
        provision_rate=Decimal("10.0000"),
    )
    client.force_login(user)
    query = {
        "branch": branch.pk,
        "start_date": (today - timedelta(days=30)).isoformat(),
        "end_date": today.isoformat(),
    }

    payloads = {
        report.slug: client.get(
            reverse(f"reports:{report.grid_route_name}"), query
        ).json()
        for report in reports
    }

    assert payloads["disbursement-register"]["records"] == 2
    assert payloads["outstanding-loans"]["userdata"]["outstanding"] == "700.00"
    assert payloads["overdue-loans"]["userdata"]["remaining"] == "110.00"
    assert payloads["aging"]["rows"][0]["bucket"] == "1-30"
    assert payloads["repayment-schedule"]["records"] == 1
    assert payloads["portfolio-at-risk"]["rows"][0]["amount"] == "700.00"
    assert payloads["portfolio-at-risk"]["rows"][0]["percentage"] == "100.00"
    assert payloads["portfolio-at-risk"]["rows"][1]["amount"] == "0.00"
    assert payloads["loan-statement"]["records"] == 1
    assert payloads["loan-statement"]["userdata"]["amount"] == "10.00"
    assert payloads["write-offs"]["userdata"]["amount"] == "300.00"
    assert payloads["classification"]["userdata"]["principal"] == "700.00"


@pytest.mark.django_db
def test_collection_reports_share_the_posted_sheet_source_and_totals(client):
    branch = Branch.objects.create(name="Collection report branch", code="RPT-COL")
    user = _create_user(branch)
    report_slugs = {
        "collection-summary",
        "collection-register",
        "collector-performance",
        "collection-efficiency",
        "shortage-excess",
    }
    reports = [report for report in REPORTS if report.slug in report_slugs]
    _grant_routes(
        user,
        tuple(
            route_name
            for report in reports
            for route_name in (report.route_name, report.grid_route_name)
        ),
    )
    center = Center.unscoped.create(
        branch=branch,
        name="Collection report center",
        code="RPT-COL-CENTER",
    )
    officer = Staff.unscoped.create(
        branch=branch,
        full_name="Collection report officer",
    )
    member = Member.unscoped.create(
        branch=branch,
        center=center,
        member_code="RPT-COL-MEMBER",
        first_name="Collection",
        last_name="Member",
        status=Member.STATUS_ACTIVE,
    )
    today = timezone.localdate()
    sheet = CollectionSheet.unscoped.create(
        branch=branch,
        center=center,
        field_officer=officer,
        collection_date=today,
        idempotency_key=uuid.uuid4(),
        generated_by=user,
        cash_received=Decimal("100.00"),
        cash_shortage=Decimal("40.00"),
        cash_excess=Decimal("0.00"),
    )
    CollectionSheetItem.unscoped.create(
        branch=branch,
        sheet=sheet,
        member=member,
        expected_installment=Decimal("100.00"),
        expected_overdue=Decimal("20.00"),
        expected_penalty=Decimal("5.00"),
        expected_savings=Decimal("10.00"),
        collected_amount=Decimal("90.00"),
        savings_amount=Decimal("5.00"),
    )
    sheet.status = CollectionSheet.STATUS_IN_PROGRESS
    sheet.save()
    sheet.status = CollectionSheet.STATUS_SUBMITTED
    sheet.save()
    sheet.status = CollectionSheet.STATUS_RECONCILED
    sheet.save()
    sheet.status = CollectionSheet.STATUS_POSTED
    sheet.posted_by = user
    sheet.posted_at = timezone.now()
    sheet.save()
    client.force_login(user)
    query = {
        "branch": branch.pk,
        "start_date": today.isoformat(),
        "end_date": today.isoformat(),
    }

    payloads = {
        report.slug: client.get(
            reverse(f"reports:{report.grid_route_name}"), query
        ).json()
        for report in reports
    }

    assert payloads["collection-summary"]["userdata"]["due"] == "135.00"
    assert payloads["collection-summary"]["userdata"]["collected"] == "95.00"
    assert payloads["collector-performance"]["userdata"]["sheets"] == 1
    assert payloads["collector-performance"]["userdata"]["due"] == "135.00"
    assert payloads["collector-performance"]["userdata"]["collected"] == "95.00"
    efficiency = Decimal(
        payloads["collection-efficiency"]["rows"][0]["efficiency"]
    )
    assert efficiency.quantize(Decimal("0.01")) == Decimal("70.37")
    assert payloads["shortage-excess"]["userdata"]["cash_received"] == "100.00"
    assert payloads["shortage-excess"]["userdata"]["shortage"] == "40.00"
