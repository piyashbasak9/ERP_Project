import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from threading import Barrier
from typing import ClassVar
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import DatabaseError, close_old_connections, connection, transaction
from django.test import RequestFactory, TestCase, TransactionTestCase
from django.urls import reverse

from core.accounting.models import (
    Account,
    AccountingPeriod,
    FiscalYear,
    SystemAccountMapping,
)
from core.accounting.services import JournalService
from core.accounts.models import Role, RolePermission, UserRole
from core.accounts.permission_utils import get_all_urls, sync_permissions
from core.branch.models import Branch, BranchClosedDay, Center, Member, Staff
from core.common.context import current_branch_id_var, current_user_var
from core.loans.models import LoanApplication, LoanInstallment, LoanPayment
from core.loans.services import LoanService
from core.savings.models import SavingsProduct, SavingsTransaction
from core.savings.services import SavingsService

from .models import (
    CollectionPosting,
    CollectionPostingReversal,
    CollectionSheet,
    CollectionSheetItem,
    CollectionSheetReversal,
)
from .services import CollectionService

ZERO = Decimal("0.00")


class CollectionFixture:
    def setUp(self):
        self.branch = Branch.unscoped.create(name="Collection Branch", code="COL")
        user_model = get_user_model()
        self.manager = user_model.objects.create_user(
            username="collection-manager", password="not-a-secret", branch=self.branch
        )
        self.collector = user_model.objects.create_user(
            username="collection-officer", password="not-a-secret", branch=self.branch
        )
        self.checker = user_model.objects.create_user(
            username="collection-checker", password="not-a-secret", branch=self.branch
        )
        self.poster = user_model.objects.create_user(
            username="collection-poster", password="not-a-secret", branch=self.branch
        )
        self.other_poster = user_model.objects.create_user(
            username="collection-poster-2", password="not-a-secret", branch=self.branch
        )
        self.current_user_token = current_user_var.set(self.manager)
        self.current_branch_token = current_branch_id_var.set(None)
        self.addCleanup(current_user_var.reset, self.current_user_token)
        self.addCleanup(current_branch_id_var.reset, self.current_branch_token)
        self.request = RequestFactory().post("/collections/sheets/")
        self.request.user = self.manager
        self.request.active_branch_id = None

        manager_staff = Staff.unscoped.create(
            branch=self.branch,
            user=self.manager,
            full_name="Branch Manager",
            designation="Branch Manager",
        )
        officer_staff = Staff.unscoped.create(
            branch=self.branch,
            user=self.collector,
            full_name="Field Officer",
            designation="Field Officer",
        )
        self.branch.branch_manager = manager_staff
        self.branch.save(update_fields=["branch_manager"])
        self.center = Center.unscoped.create(
            branch=self.branch,
            field_officer=officer_staff,
            name="Collection Center",
            code="COL-CENTER",
        )
        self.officer_staff = officer_staff
        self.member = self.create_member("COL-0001")

        fiscal_year = FiscalYear.objects.create(
            name="Collections FY 2026",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
        )
        AccountingPeriod.objects.create(
            fiscal_year=fiscal_year,
            number=1,
            name="FY 2026",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
        )
        self._create_account_mappings()
        self.loan_product = self._loan_product()
        self.savings_product = SavingsProduct.objects.create(
            code="COLSAV",
            name="Collection savings",
            savings_type=SavingsProduct.TYPE_GENERAL,
            minimum_opening_balance=ZERO,
        )
        self.loan = self.create_loan(self.member)
        self.installment = LoanInstallment.objects.get(
            loan=self.loan,
            schedule_version=self.loan.schedule_version,
            number=1,
        )

    def service_as(self, user):
        current_user_var.set(user)
        self.request.user = user
        return CollectionService(self.request)

    def loan_service_as(self, user):
        current_user_var.set(user)
        self.request.user = user
        return LoanService(self.request)

    def savings_service_as(self, user):
        current_user_var.set(user)
        self.request.user = user
        return SavingsService(self.request)

    def create_member(self, code, *, branch=None, center=None):
        return Member.unscoped.create(
            branch=branch or self.branch,
            center=center or self.center,
            member_code=code,
            first_name="Collection",
            last_name=code,
            status=Member.STATUS_ACTIVE,
        )

    def create_loan(self, member):
        maker_service = self.loan_service_as(self.manager)
        application = maker_service.create_application(
            branch_id=member.branch_id,
            member_id=member.pk,
            product_id=self.loan_product.pk,
            amount=Decimal("1200.00"),
            term=12,
            purpose="Collection test loan",
            idempotency_key=uuid.uuid4(),
        )
        maker_service.submit(application_id=application.pk)
        self.loan_service_as(self.checker).review(
            application_id=application.pk,
            decision=LoanApplication.STATUS_APPROVED,
            comment="Approved for collection test.",
            approved_amount=Decimal("1200.00"),
        )
        loan = self.loan_service_as(self.poster).disburse(
            application_id=application.pk,
            disbursement_date=date(2026, 1, 1),
            channel="cash_in_hand",
            idempotency_key=uuid.uuid4(),
        )
        return loan

    def open_savings(self, member):
        return self.savings_service_as(self.manager).open_account(
            branch_id=member.branch_id,
            product_id=self.savings_product.pk,
            opened_at=date(2026, 1, 1),
            member_id=member.pk,
            opening_balance=ZERO,
            idempotency_key=uuid.uuid4(),
        )

    def prepare_sheet(
        self,
        *,
        amount=None,
        savings_amount=ZERO,
        cash_received=None,
        collection_date=None,
        member=None,
        group_id=None,
    ):
        collection_date = collection_date or self.installment.due_date
        member = member or self.member
        sheet = self.service_as(self.manager).generate(
            center_id=self.center.pk,
            group_id=group_id,
            collection_date=collection_date,
            idempotency_key=uuid.uuid4(),
        )
        item = CollectionSheetItem.unscoped.get(sheet=sheet, member=member)
        amount = item.expected_installment + item.expected_overdue if amount is None else amount
        self.service_as(self.collector).start(sheet_id=sheet.pk)
        self.service_as(self.collector).record_collection(
            sheet_id=sheet.pk,
            item_id=item.pk,
            collected_amount=amount,
            savings_amount=savings_amount,
            reason="Part payment" if amount < item.expected_installment + item.expected_overdue else "",
        )
        self.service_as(self.collector).submit(sheet_id=sheet.pk)
        cash_received = (
            amount + savings_amount
            if cash_received is None
            else cash_received
        )
        sheet = self.service_as(self.checker).reconcile(
            sheet_id=sheet.pk, cash_received=cash_received
        )
        return sheet, item

    def _loan_product(self):
        from core.loans.models import LoanProduct

        return LoanProduct.objects.create(
            code="COLLOAN",
            name="Collection test product",
            interest_method=LoanProduct.METHOD_FLAT,
            annual_interest_rate=Decimal("12.0000"),
            min_amount=Decimal("100.00"),
            max_amount=Decimal("10000.00"),
            min_term=1,
            max_term=24,
            installment_frequency=LoanProduct.FREQUENCY_MONTHLY,
            requires_verified_kyc=False,
        )

    def _create_account_mappings(self):
        account_types = {
            "cash_in_hand": (Account.ASSET, Account.DEBIT),
            "bank_accounts": (Account.ASSET, Account.DEBIT),
            "loans_to_members": (Account.ASSET, Account.DEBIT),
            "loan_loss_provision": (Account.ASSET, Account.CREDIT),
            "member_savings": (Account.LIABILITY, Account.CREDIT),
            "interest_income": (Account.INCOME, Account.CREDIT),
            "interest_expense": (Account.EXPENSE, Account.DEBIT),
            "fee_income": (Account.INCOME, Account.CREDIT),
            "insurance_income": (Account.INCOME, Account.CREDIT),
            "penalty_income": (Account.INCOME, Account.CREDIT),
            "loan_overpayment_payable": (Account.LIABILITY, Account.CREDIT),
            "recovery_income": (Account.INCOME, Account.CREDIT),
            "withholding_tax_payable": (Account.LIABILITY, Account.CREDIT),
            "cash_shortage": (Account.EXPENSE, Account.DEBIT),
            "collection_excess_income": (Account.INCOME, Account.CREDIT),
        }
        for index, (key, (account_type, normal_balance)) in enumerate(
            account_types.items(), start=1
        ):
            account = Account.objects.create(
                code=f"CC{index:03d}",
                name=key.replace("_", " ").title(),
                account_type=account_type,
                normal_balance=normal_balance,
            )
            SystemAccountMapping.objects.create(key=key, account=account)


class CollectionWorkflowTests(CollectionFixture, TransactionTestCase):
    def test_generation_is_idempotent_and_includes_expected_installment(self):
        key = uuid.uuid4()
        service = self.service_as(self.manager)
        sheet = service.generate(
            center_id=self.center.pk,
            collection_date=self.installment.due_date,
            idempotency_key=key,
        )
        replay = service.generate(
            center_id=self.center.pk,
            collection_date=self.installment.due_date,
            idempotency_key=key,
        )
        item = CollectionSheetItem.unscoped.get(sheet=sheet, member=self.member)

        self.assertEqual(replay.pk, sheet.pk)
        self.assertEqual(item.expected_installment, self.installment.total_due)
        self.assertEqual(item.expected_overdue, ZERO)
        self.assertEqual(
            CollectionSheet.unscoped.filter(
                center=self.center, collection_date=self.installment.due_date
            ).count(),
            1,
        )

    def test_partial_collection_requires_reason_and_posts_allocation(self):
        amount = self.installment.total_due / 2
        sheet, item = self.prepare_sheet(amount=amount)
        posting_key = uuid.uuid4()
        payment_key = uuid.uuid5(
            posting_key,
            f"collection-{sheet.pk}-{item.pk}-loan-{self.loan.pk}",
        )
        posted = self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=posting_key
        )
        item.refresh_from_db()
        payment = LoanPayment.objects.get(loan=self.loan, idempotency_key=payment_key)

        self.assertEqual(posted.status, CollectionSheet.STATUS_POSTED)
        self.assertEqual(item.reason, "Part payment")
        self.assertEqual(payment.amount, amount)
        self.assertEqual(
            payment.amount,
            payment.penalty_paid + payment.fee_paid + payment.interest_paid
            + payment.principal_paid + payment.overpayment_amount,
        )

    def test_full_and_advance_payments_use_existing_loan_allocation(self):
        sheet, _ = self.prepare_sheet(amount=self.installment.total_due)
        full_payment = self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        self.assertEqual(full_payment.status, CollectionSheet.STATUS_POSTED)
        payment = LoanPayment.objects.get(loan=self.loan)
        self.assertEqual(payment.amount, self.installment.total_due)

        member2 = self.create_member("COL-0002")
        loan2 = self.create_loan(member2)
        first_due = LoanInstallment.objects.get(
            loan=loan2, schedule_version=loan2.schedule_version, number=1
        )
        second_due = LoanInstallment.objects.get(
            loan=loan2, schedule_version=loan2.schedule_version, number=2
        )
        third_due = LoanInstallment.objects.get(
            loan=loan2, schedule_version=loan2.schedule_version, number=3
        )
        sheet2, _ = self.prepare_sheet(
            member=member2,
            amount=first_due.total_due + second_due.total_due + third_due.total_due,
            collection_date=second_due.due_date,
        )
        self.service_as(self.poster).post(
            sheet_id=sheet2.pk, idempotency_key=uuid.uuid4()
        )
        advance_payment = LoanPayment.objects.get(loan=loan2)

        self.assertGreater(advance_payment.advance_amount, ZERO)
        third_due.refresh_from_db()
        self.assertGreater(third_due.principal_paid, ZERO)

    def test_reconciliation_variances_post_to_mapped_accounts(self):
        total = self.installment.total_due
        shortage_sheet, _ = self.prepare_sheet(
            amount=total,
            cash_received=total - Decimal("5.00"),
        )
        self.service_as(self.poster).post(
            sheet_id=shortage_sheet.pk, idempotency_key=uuid.uuid4()
        )
        shortage_sheet.refresh_from_db()
        shortage_lines = list(shortage_sheet.variance_journal.lines.select_related("account"))
        self.assertEqual(shortage_sheet.cash_shortage, Decimal("5.00"))
        self.assertEqual(
            sum(line.debit for line in shortage_lines if line.account.code == "CC014"),
            Decimal("5.00"),
        )
        self.assertEqual(
            sum(line.credit for line in shortage_lines if line.account.code == "CC001"),
            Decimal("5.00"),
        )

        member2 = self.create_member("COL-0002")
        loan2 = self.create_loan(member2)
        installment2 = LoanInstallment.objects.get(
            loan=loan2, schedule_version=loan2.schedule_version, number=1
        )
        excess_sheet, _ = self.prepare_sheet(
            member=member2,
            amount=installment2.total_due,
            cash_received=installment2.total_due + Decimal("7.00"),
            collection_date=installment2.due_date + timedelta(days=1),
        )
        self.service_as(self.poster).post(
            sheet_id=excess_sheet.pk, idempotency_key=uuid.uuid4()
        )
        excess_sheet.refresh_from_db()
        self.assertEqual(excess_sheet.cash_excess, Decimal("7.00"))
        excess_lines = list(excess_sheet.variance_journal.lines.select_related("account"))
        self.assertEqual(
            sum(line.debit for line in excess_lines if line.account.code == "CC001"),
            Decimal("7.00"),
        )
        self.assertEqual(
            sum(line.credit for line in excess_lines if line.account.code == "CC015"),
            Decimal("7.00"),
        )

    def test_savings_deposit_and_loan_journals_equal_the_sheet_total(self):
        account = self.open_savings(self.member)
        amount = self.installment.total_due
        sheet, item = self.prepare_sheet(
            amount=amount, savings_amount=Decimal("25.00")
        )
        self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        item.refresh_from_db()
        postings = list(
            CollectionPosting.objects.filter(item=item).select_related(
                "loan_payment__journal_entry", "savings_transaction__journal_entry"
            )
        )
        posted_total = sum((posting.amount for posting in postings), ZERO)
        journal_lines = [
            line
            for posting in postings
            for line in (
                posting.loan_payment.journal_entry.lines.all()
                if posting.loan_payment_id
                else posting.savings_transaction.journal_entry.lines.all()
            )
        ]
        journal_debits = sum((line.debit for line in journal_lines), ZERO)
        journal_credits = sum((line.credit for line in journal_lines), ZERO)

        self.assertEqual(account.pk, postings[-1].savings_transaction.account_id)
        self.assertEqual(posted_total, amount + Decimal("25.00"))
        self.assertEqual(journal_debits, posted_total)
        self.assertEqual(journal_credits, posted_total)
        self.assertEqual(journal_debits, journal_credits)
        self.assertEqual(
            sum(posting.loan_payment.amount for posting in postings if posting.loan_payment_id)
            + sum(
                posting.savings_transaction.amount
                for posting in postings
                if posting.savings_transaction_id
            ),
            item.collected_amount + item.savings_amount,
        )

    def test_failed_row_rolls_back_every_sheet_posting(self):
        member2 = self.create_member("COL-0002")
        loan2 = self.create_loan(member2)
        account1 = self.open_savings(self.member)
        account2 = self.open_savings(member2)
        sheet = self.service_as(self.manager).generate(
            center_id=self.center.pk,
            collection_date=self.installment.due_date,
            idempotency_key=uuid.uuid4(),
        )
        items = list(
            CollectionSheetItem.unscoped.filter(sheet=sheet).order_by("member__member_code")
        )
        self.service_as(self.collector).start(sheet_id=sheet.pk)
        for item in items:
            loan = self.loan if item.member_id == self.member.pk else loan2
            due = LoanInstallment.objects.get(
                loan=loan, schedule_version=loan.schedule_version, number=1
            )
            self.service_as(self.collector).record_collection(
                sheet_id=sheet.pk,
                item_id=item.pk,
                collected_amount=due.total_due,
                savings_amount=Decimal("10.00"),
            )
        self.service_as(self.collector).submit(sheet_id=sheet.pk)
        self.service_as(self.checker).reconcile(
            sheet_id=sheet.pk,
            cash_received=sum((item.collected_amount + item.savings_amount for item in items), ZERO),
        )
        original_deposit = SavingsService.deposit
        calls = 0

        def fail_second_deposit(service, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise ValidationError("Injected row failure")
            return original_deposit(service, **kwargs)

        with patch(
            "core.collections.services.SavingsService.deposit",
            new=fail_second_deposit,
        ), self.assertRaises(ValidationError):
            self.service_as(self.poster).post(
                sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
            )

        sheet.refresh_from_db()
        self.assertEqual(sheet.status, CollectionSheet.STATUS_RECONCILED)
        self.assertFalse(LoanPayment.objects.filter(loan__in=[self.loan, loan2]).exists())
        self.assertFalse(SavingsTransaction.objects.filter(account__in=[account1, account2]).exists())
        self.assertFalse(CollectionPosting.objects.filter(item__sheet=sheet).exists())
        self.assertIsNone(sheet.variance_journal_id)

    def test_lifecycle_rejects_bad_transitions_and_collector_cannot_verify(self):
        sheet = self.service_as(self.manager).generate(
            center_id=self.center.pk,
            collection_date=date(2026, 1, 15),
            idempotency_key=uuid.uuid4(),
        )
        outsider = get_user_model().objects.create_user(
            username="collection-outsider",
            branch=self.branch,
        )
        with self.assertRaises(PermissionDenied):
            self.service_as(outsider).start(sheet_id=sheet.pk)
        with self.assertRaises(ValidationError):
            self.service_as(self.collector).submit(sheet_id=sheet.pk)
        self.service_as(self.collector).start(sheet_id=sheet.pk)
        with self.assertRaises(ValidationError):
            self.service_as(self.collector).start(sheet_id=sheet.pk)
        self.service_as(self.collector).submit(sheet_id=sheet.pk)
        with self.assertRaises(ValidationError):
            self.service_as(self.collector).reconcile(
                sheet_id=sheet.pk, cash_received=ZERO
            )
        self.service_as(self.checker).reconcile(
            sheet_id=sheet.pk, cash_received=ZERO
        )
        with self.assertRaises(ValidationError):
            self.service_as(self.collector).start(sheet_id=sheet.pk)

    def test_two_verifiers_and_same_key_post_a_sheet_only_once(self):
        sheet, _ = self.prepare_sheet(
            amount=ZERO,
            collection_date=date(2026, 1, 15),
            cash_received=ZERO,
        )
        key = uuid.uuid4()
        barrier = Barrier(2)

        def post_as(user):
            close_old_connections()
            request = RequestFactory().post("/collections/")
            request.user = get_user_model().objects.get(pk=user.pk)
            request.active_branch_id = None
            barrier.wait(timeout=10)
            result = CollectionService(request).post(
                sheet_id=sheet.pk, idempotency_key=key
            )
            close_old_connections()
            return result.status

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(
                executor.map(post_as, [self.poster, self.other_poster])
            )
        replay = self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=key
        )

        self.assertEqual(results, [CollectionSheet.STATUS_POSTED] * 2)
        self.assertEqual(replay.status, CollectionSheet.STATUS_POSTED)
        self.assertEqual(CollectionPosting.objects.filter(item__sheet=sheet).count(), 0)
        self.assertEqual(
            CollectionSheet.unscoped.filter(pk=sheet.pk, status="posted").count(), 1
        )

    def test_collection_sheet_reversal_restores_loan_and_nets_journal_once(self):
        original_outstanding = self.loan.outstanding_principal
        savings_account = self.open_savings(self.member)
        sheet, _ = self.prepare_sheet(
            amount=self.installment.total_due,
            savings_amount=Decimal("25.00"),
        )
        self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        payment = LoanPayment.objects.get(loan=self.loan)
        self.loan.refresh_from_db()
        savings_account.refresh_from_db()
        outstanding_after_post = self.loan.outstanding_principal
        savings_after_post = savings_account.balance
        reversal_date = payment.payment_date + timedelta(days=1)
        reversal = self.service_as(self.manager).request_reversal(
            sheet_id=sheet.pk,
            reversal_date=reversal_date,
            reason="Correct duplicated collection sheet",
            idempotency_key=uuid.uuid4(),
        )
        with self.assertRaises(ValidationError):
            self.service_as(self.manager).approve_reversal(reversal_id=reversal.pk)

        approved = self.service_as(self.checker).approve_reversal(
            reversal_id=reversal.pk
        )
        self.loan.refresh_from_db()
        savings_account.refresh_from_db()
        sheet.refresh_from_db()
        payment.refresh_from_db()
        original_transactions = list(
            CollectionPosting.objects.filter(item__sheet=sheet)
        )
        original_lines = [
            line
            for posting in original_transactions
            for line in (
                posting.loan_payment.journal_entry.lines.all()
                if posting.loan_payment_id
                else posting.savings_transaction.journal_entry.lines.all()
            )
        ]
        reversed_lines = [
            line
            for posting_reversal in approved.posting_reversals.all()
            for line in (
                posting_reversal.loan_payment_reversal.journal_entry.lines.all()
                if posting_reversal.loan_payment_reversal_id
                else posting_reversal.savings_transaction_reversal.journal_entry.lines.all()
            )
        ]
        balances = {}
        for line in original_lines + reversed_lines:
            account_id, debit, credit = line.account_id, line.debit, line.credit
            balances[account_id] = balances.get(account_id, ZERO) + debit - credit

        self.assertEqual(approved.status, CollectionSheetReversal.STATUS_COMPLETED)
        self.assertEqual(sheet.status, CollectionSheet.STATUS_REVERSED)
        self.assertEqual(self.loan.outstanding_principal, original_outstanding)
        self.assertEqual(savings_account.balance, ZERO)
        self.assertEqual(approved.approved_by_id, self.checker.pk)
        self.assertTrue(all(balance == ZERO for balance in balances.values()))
        replay = self.service_as(self.other_poster).approve_reversal(
            reversal_id=reversal.pk
        )
        self.assertEqual(replay.pk, approved.pk)
        self.assertEqual(
            CollectionPostingReversal.objects.filter(
                sheet_reversal=approved
            ).count(),
            2,
        )
        self.assertLess(outstanding_after_post, original_outstanding)
        self.assertGreater(savings_after_post, ZERO)

    def test_collection_sheet_reversal_rejects_closed_accounting_period(self):
        sheet, _ = self.prepare_sheet(amount=self.installment.total_due)
        self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        payment = LoanPayment.objects.get(loan=self.loan)
        reversal = self.service_as(self.manager).request_reversal(
            sheet_id=sheet.pk,
            reversal_date=payment.payment_date + timedelta(days=1),
            reason="Closed-period reversal test",
            idempotency_key=uuid.uuid4(),
        )
        AccountingPeriod.objects.update(is_closed=True)
        with self.assertRaisesRegex(ValidationError, "open accounting period"):
            self.service_as(self.checker).approve_reversal(reversal_id=reversal.pk)
        sheet.refresh_from_db()
        self.assertEqual(sheet.status, CollectionSheet.STATUS_POSTED)
        self.assertFalse(
            CollectionPostingReversal.objects.filter(
                sheet_reversal=reversal
            ).exists()
        )

    def test_collection_sheet_reversal_is_safe_under_concurrent_approval(self):
        sheet, _ = self.prepare_sheet(amount=self.installment.total_due)
        self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        payment = LoanPayment.objects.get(loan=self.loan)
        reversal = self.service_as(self.manager).request_reversal(
            sheet_id=sheet.pk,
            reversal_date=payment.payment_date + timedelta(days=1),
            reason="Concurrent reversal test",
            idempotency_key=uuid.uuid4(),
        )
        barrier = Barrier(2)

        def approve_as(user):
            close_old_connections()
            user_token = current_user_var.set(
                get_user_model().objects.get(pk=user.pk)
            )
            branch_token = current_branch_id_var.set(None)
            request = RequestFactory().post("/collections/reversals/")
            request.user = current_user_var.get()
            request.active_branch_id = None
            barrier.wait(timeout=10)
            result = CollectionService(request).approve_reversal(
                reversal_id=reversal.pk
            )
            current_branch_id_var.reset(branch_token)
            current_user_var.reset(user_token)
            close_old_connections()
            return result.status

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(
                executor.map(approve_as, [self.checker, self.other_poster])
            )

        self.assertEqual(
            results,
            [CollectionSheetReversal.STATUS_COMPLETED] * 2,
        )
        self.assertEqual(
            CollectionPostingReversal.objects.filter(
                sheet_reversal=reversal
            ).count(),
            1,
        )

    def test_closed_business_date_rejects_all_financial_posting_services(self):
        savings_account = self.open_savings(self.member)
        sheet, _ = self.prepare_sheet(
            amount=self.installment.total_due,
            collection_date=self.installment.due_date,
        )
        BranchClosedDay.objects.create(
            branch=self.branch,
            business_date=self.installment.due_date,
            cash_count=ZERO,
            closed_by=self.manager,
        )
        closed_message = r"Business date .* is closed"

        with self.assertRaisesRegex(ValidationError, closed_message):
            self.service_as(self.poster).post(
                sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
            )
        with self.assertRaisesRegex(ValidationError, closed_message):
            self.loan_service_as(self.poster).repay(
                loan_id=self.loan.pk,
                amount=Decimal("1.00"),
                payment_date=self.installment.due_date,
                idempotency_key=uuid.uuid4(),
            )
        with self.assertRaisesRegex(ValidationError, closed_message):
            self.savings_service_as(self.manager).deposit(
                account_id=savings_account.pk,
                amount=Decimal("1.00"),
                transaction_date=self.installment.due_date,
                payment_channel="cash",
                idempotency_key=uuid.uuid4(),
            )
        with self.assertRaisesRegex(ValidationError, closed_message):
            JournalService(self.request).create_entry(
                branch_id=self.branch.pk,
                journal_date=self.installment.due_date,
                narration="Closed-date test",
            )

    def test_runtime_role_cannot_mutate_posted_sheets_or_rows(self):
        sheet, item = self.prepare_sheet(
            amount=self.installment.total_due,
            collection_date=date(2026, 1, 15),
        )
        self.service_as(self.poster).post(
            sheet_id=sheet.pk, idempotency_key=uuid.uuid4()
        )
        payment = LoanPayment.objects.get(loan=self.loan)
        reversal = self.service_as(self.manager).request_reversal(
            sheet_id=sheet.pk,
            reversal_date=payment.payment_date + timedelta(days=1),
            reason="Runtime immutability test",
            idempotency_key=uuid.uuid4(),
        )
        self.service_as(self.checker).approve_reversal(reversal_id=reversal.pk)
        posting_reversal = CollectionPostingReversal.objects.get(
            sheet_reversal=reversal
        )
        payment_reversal = payment.reversal
        approval = reversal.approval_record
        database_dir = Path(__file__).resolve().parents[2] / "database"

        with connection.cursor() as cursor:
            cursor.execute("GRANT USAGE ON SCHEMA public TO erp_runtime")
            cursor.execute(
                "GRANT SELECT, INSERT ON ALL TABLES IN SCHEMA public TO erp_runtime"
            )
            cursor.execute(
                "GRANT USAGE, SELECT ON ALL SEQUENCES IN SCHEMA public TO erp_runtime"
            )
            cursor.execute(
                (database_dir / "savings_runtime_role.sql").read_text(
                    encoding="utf-8"
                )
            )
            cursor.execute("SET ROLE erp_runtime")
            try:
                for statement, params in (
                    (
                        "UPDATE collection_sheets SET cash_received = 1 WHERE id = %s",
                        [sheet.pk],
                    ),
                    (
                        "UPDATE collection_sheet_items SET collected_amount = 1 WHERE id = %s",
                        [item.pk],
                    ),
                    ("DELETE FROM collection_sheet_items WHERE id = %s", [item.pk]),
                    (
                        "UPDATE collection_sheet_reversals SET reason = 'tampered' WHERE id = %s",
                        [reversal.pk],
                    ),
                    (
                        "DELETE FROM collection_sheet_reversals WHERE id = %s",
                        [reversal.pk],
                    ),
                    (
                        "UPDATE collection_sheet_reversal_approvals SET completed_at = now() WHERE id = %s",
                        [approval.pk],
                    ),
                    (
                        "DELETE FROM collection_sheet_reversal_approvals WHERE id = %s",
                        [approval.pk],
                    ),
                    (
                        "UPDATE collection_posting_reversals SET created_at = now() WHERE id = %s",
                        [posting_reversal.pk],
                    ),
                    (
                        "DELETE FROM collection_posting_reversals WHERE id = %s",
                        [posting_reversal.pk],
                    ),
                    (
                        "UPDATE loan_payment_reversals SET reason = 'tampered' WHERE id = %s",
                        [payment_reversal.pk],
                    ),
                    (
                        "DELETE FROM loan_payment_reversals WHERE id = %s",
                        [payment_reversal.pk],
                    ),
                ):
                    with (
                        self.subTest(statement=statement),
                        self.assertRaises(DatabaseError),
                        transaction.atomic(),
                    ):
                        cursor.execute(statement, params)
            finally:
                cursor.execute("RESET ROLE")


class CollectionRouteTests(CollectionFixture, TestCase):
    route_args: ClassVar[dict[str, list[int]]] = {
        "Module_Collections_sheet_detail": [1],
        "Module_Collections_sheet_start": [1],
        "Module_Collections_sheet_entry": [1, 1],
        "Module_Collections_sheet_submit": [1],
        "Module_Collections_sheet_cancel": [1],
        "Module_Collections_sheet_print": [1],
        "Module_Collections_sheet_reconcile": [1],
        "Module_Collections_sheet_post": [1],
        "Module_Collections_sheet_reversal_request": [1],
        "Module_Collections_sheet_reversal_approve": [1],
    }

    def setUp(self):
        super().setUp()
        sync_permissions()
        self.role = Role.objects.create(name="collections-test-role")
        UserRole.objects.create(user=self.manager, role=self.role)
        self.client.force_login(self.manager)

    def test_every_collection_route_requires_its_explicit_permission(self):
        self.client.logout()
        self.client.force_login(self.other_poster)
        role = Role.objects.create(name="collections-denied-role")
        UserRole.objects.create(user=self.other_poster, role=role)
        route_names = {
            route["route_name"]
            for route in get_all_urls()
            if route["route_name"].startswith("Module_Collections_")
        }
        for route_name in sorted(route_names):
            with self.subTest(route_name=route_name):
                response = self.client.get(
                    reverse(
                        f"collections:{route_name}",
                        args=self.route_args.get(route_name, []),
                    )
                )
                self.assertEqual(response.status_code, 403)

    def test_state_changing_routes_reject_get(self):
        state_routes = [
            "Module_Collections_sheet_start",
            "Module_Collections_sheet_entry",
            "Module_Collections_sheet_submit",
            "Module_Collections_sheet_cancel",
            "Module_Collections_sheet_reconcile",
            "Module_Collections_sheet_post",
            "Module_Collections_sheet_reversal_request",
            "Module_Collections_sheet_reversal_approve",
        ]
        from core.accounts.models import Permission

        for route_name in state_routes:
            permission = Permission.objects.get(route_name=route_name)
            RolePermission.objects.create(role=self.role, permission=permission)
        for route_name in state_routes:
            with self.subTest(route_name=route_name):
                response = self.client.get(
                    reverse(
                        f"collections:{route_name}",
                        args=self.route_args[route_name],
                    )
                )
                self.assertEqual(response.status_code, 405)

    def test_collection_entry_screen_renders_live_total_and_unsaved_warning(self):
        from core.accounts.models import Permission

        permission = Permission.objects.get(
            route_name="Module_Collections_sheet_detail"
        )
        RolePermission.objects.create(role=self.role, permission=permission)
        sheet = self.service_as(self.manager).generate(
            center_id=self.center.pk,
            collection_date=self.installment.due_date,
            idempotency_key=uuid.uuid4(),
        )
        self.service_as(self.manager).start(sheet_id=sheet.pk)

        response = self.client.get(
            reverse(
                "collections:Module_Collections_sheet_detail",
                args=[sheet.pk],
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="table-responsive"')
        self.assertContains(response, 'id="live-total"')
        self.assertContains(response, 'id="unsaved-warning"')
        self.assertContains(response, 'data-original="0.00"')

    def test_forged_branch_ids_cannot_open_enter_submit_reconcile_post_or_export(self):
        branch_b = Branch.unscoped.create(name="Other Branch", code="COL-B")
        user_b = get_user_model().objects.create_user(
            username="branch-b-officer",
            password="not-a-secret",
            branch=branch_b,
        )
        manager_b = Staff.unscoped.create(
            branch=branch_b, user=user_b, full_name="Other Manager"
        )
        officer_b = Staff.unscoped.create(
            branch=branch_b,
            user=self.other_poster,
            full_name="Other Officer",
        )
        center_b = Center.unscoped.create(
            branch=branch_b,
            field_officer=officer_b,
            name="Private Branch B Center",
            code="COL-B-CENTER",
        )
        member_b = Member.unscoped.create(
            branch=branch_b,
            center=center_b,
            member_code="COL-B-0001",
            first_name="Private",
            last_name="Member",
            status=Member.STATUS_ACTIVE,
        )
        sheet_b = CollectionSheet.unscoped.create(
            branch=branch_b,
            center=center_b,
            field_officer=officer_b,
            collection_date=date(2026, 1, 15),
            idempotency_key=uuid.uuid4(),
            generated_by=user_b,
        )
        item_b = CollectionSheetItem.unscoped.create(
            branch=branch_b,
            sheet=sheet_b,
            member=member_b,
        )
        with connection.cursor() as cursor:
            for status in (
                CollectionSheet.STATUS_IN_PROGRESS,
                CollectionSheet.STATUS_SUBMITTED,
                CollectionSheet.STATUS_RECONCILED,
                CollectionSheet.STATUS_POSTED,
            ):
                cursor.execute(
                    "UPDATE collection_sheets SET status = %s WHERE id = %s",
                    [status, sheet_b.pk],
                )
        reversal_b = CollectionSheetReversal.objects.create(
            sheet=sheet_b,
            reversal_date=date(2026, 1, 16),
            idempotency_key=uuid.uuid4(),
            reason="Forged branch reversal",
            requested_by=user_b,
        )
        del manager_b
        from core.accounts.models import Permission

        route_names = {
            route["route_name"]
            for route in get_all_urls()
            if route["route_name"].startswith("Module_Collections_")
        }
        RolePermission.objects.bulk_create(
            [
                RolePermission(role=self.role, permission=permission)
                for permission in Permission.objects.filter(route_name__in=route_names)
                if not self.role.permissions.filter(pk=permission.pk).exists()
            ]
        )
        response = self.client.get(reverse("collections:Module_Collections_sheet_list"))
        self.assertNotContains(response, "Private Branch B Center")
        self.assertEqual(
            self.client.get(
                reverse("collections:Module_Collections_sheet_detail", args=[sheet_b.pk])
            ).status_code,
            404,
        )
        post_routes = [
            ("Module_Collections_sheet_start", [sheet_b.pk], {}),
            ("Module_Collections_sheet_entry", [sheet_b.pk, item_b.pk], {
                "collected_amount": "0.00",
                "savings_amount": "0.00",
                "reason": "",
            }),
            ("Module_Collections_sheet_submit", [sheet_b.pk], {}),
            ("Module_Collections_sheet_reconcile", [sheet_b.pk], {"cash_received": "0.00"}),
            ("Module_Collections_sheet_post", [sheet_b.pk], {"idempotency_key": str(uuid.uuid4())}),
            (
                "Module_Collections_sheet_reversal_request",
                [sheet_b.pk],
                {
                    "reversal_date": "2026-01-16",
                    "reason": "Forged branch reversal",
                    "idempotency_key": str(uuid.uuid4()),
                },
            ),
            ("Module_Collections_sheet_reversal_approve", [reversal_b.pk], {}),
        ]
        for route_name, args, data in post_routes:
            with self.subTest(route_name=route_name):
                response = self.client.post(
                    reverse(f"collections:{route_name}", args=args), data
                )
                self.assertEqual(response.status_code, 403)

        excel = self.client.get(reverse("collections:Module_Collections_report_excel"))
        self.assertEqual(excel.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(excel.content), read_only=True)
        values = [cell for row in workbook.active.iter_rows(values_only=True) for cell in row]
        self.assertNotIn("Private Branch B Center", values)
        pdf = self.client.get(reverse("collections:Module_Collections_report_pdf"))
        self.assertEqual(pdf.status_code, 200)
        self.assertNotIn(b"COL-B-CENTER", pdf.content)
